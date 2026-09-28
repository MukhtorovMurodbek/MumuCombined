"""ConvertBot's formats, converters, the worker process that runs them, and the queue.

Every section below, from its `# ─── module:` line to the next, is a module of its own: main.py
loads each one separately, for each bot that uses it, so a bot's `import db` is still that
bot's db. This file is never imported whole -- run main.py.
"""
raise ImportError("convert.py is loaded section by section by main.py -- run main.py")


# ─── module: convert_bot.formats ─────────────────────────────────────────────
"""What this bot converts, what it converts it to, and what this particular host can actually do."""
from __future__ import annotations

import importlib.util
import logging
import os
import shutil
import subprocess

logger = logging.getLogger(__name__)

IMAGE = "image"
VIDEO = "video"
AUDIO = "audio"
DOCUMENT = "document"
DATA = "data"
SUBTITLE = "subtitle"

CATEGORIES = (IMAGE, VIDEO, AUDIO, DOCUMENT, DATA, SUBTITLE)

ALIASES = {
    "jpeg": "jpg", "jpe": "jpg", "jfif": "jpg",
    "tif": "tiff",
    "dib": "bmp",
    "j2c": "jp2", "jpf": "jp2", "jpx": "jp2",
    "pgm": "ppm", "pbm": "ppm", "pnm": "ppm",
    "yml": "yaml",
    "markdown": "md", "mdown": "md",
    "htm": "html", "xhtml": "html",
    "ssa": "ass",
    "oga": "ogg",
    "mpeg": "mpg", "m2v": "mpg",
    "mts": "m2ts",
    "aif": "aiff", "aifc": "aiff",
    "text": "txt",
}

SOURCES = {
    IMAGE: (
        "png", "jpg", "webp", "bmp", "tiff", "ico", "cur",
        "heic", "heif", "avif", "jp2", "psd", "tga", "pcx", "dds",
        "ppm", "sgi", "xbm", "xpm", "icns", "qoi", "blp", "im", "svg",
    ),
    VIDEO: (
        "mp4", "mov", "m4v", "avi", "webm", "mkv", "flv", "wmv",
        "ts", "m2ts", "mpg", "3gp", "ogv", "asf", "gif", "apng",
    ),
    AUDIO: (
        "mp3", "wav", "ogg", "opus", "m4a", "aac", "flac", "aiff",
        "wma", "amr", "ac3", "mp2", "au", "caf", "w64", "voc",
    ),
    DOCUMENT: (
        "pdf", "docx", "epub", "mobi", "fb2", "cbz", "xps", "md",
        "txt", "html",
    ),
    DATA: ("csv", "tsv", "json", "yaml", "xlsx", "xml"),
    SUBTITLE: ("srt", "vtt", "ass"),
}

CATEGORY_OF: dict[str, str] = {}
for _category, _exts in SOURCES.items():
    for _ext in _exts:
        CATEGORY_OF.setdefault(_ext, _category)

ANIMATED_CAPABLE = ("gif", "webp", "apng")
GIF_IS_VIDEO = True

ANIMATED_TARGETS = ("mp4", "webm", "gif", "webp", "apng", "mkv", "mov", "png", "jpg")

TARGETS_BY_CATEGORY = {
    IMAGE: (
        "png", "jpg", "webp", "pdf", "ico", "tiff",
        "bmp", "gif", "avif", "heic", "jp2", "tga", "ppm", "qoi", "icns",
    ),
    VIDEO: (
        "mp4", "webm", "gif", "mp3", "mkv", "mov",
        "avi", "webp", "apng", "m4a", "aac", "ogg", "opus", "wav", "flac",
    ),
    AUDIO: (
        "mp3", "m4a", "ogg", "wav", "flac", "opus",
        "aac", "aiff", "wma", "ac3",
    ),
    DATA: (
        "json", "csv", "xlsx", "yaml", "md", "html",
        "tsv", "pdf", "txt",
    ),
    SUBTITLE: ("srt", "vtt", "ass", "txt"),

}

TARGETS_BY_FORMAT = {
    "pdf": ("docx", "png", "jpg", "txt", "html", "webp", "svg", "tiff"),
    "epub": ("pdf", "docx", "txt", "html", "png", "jpg"),
    "mobi": ("pdf", "docx", "txt", "html", "png", "jpg"),
    "fb2": ("pdf", "docx", "txt", "html", "png", "jpg"),
    "xps": ("pdf", "docx", "txt", "html", "png", "jpg"),
    "cbz": ("pdf", "png", "jpg"),
    "docx": ("pdf", "html", "md", "txt"),
    "md": ("pdf", "docx", "html", "txt"),
    "html": ("pdf", "docx", "txt", "png", "jpg"),
    "txt": ("pdf", "docx", "html"),

    "svg": ("png", "pdf", "jpg", "webp"),

    "gif": ("mp4", "webm", "webp", "apng", "png", "jpg"),
    "apng": ("gif", "mp4", "webm", "webp", "png"),

    "flac": ("wav", "aiff", "mp3", "m4a", "ogg", "opus", "aac"),
    "wav": ("mp3", "flac", "m4a", "ogg", "opus", "aiff", "aac"),
    "xml": ("json", "yaml", "md", "html", "txt", "pdf"),
}

POPULAR_COUNT = 6

FORMAT_NEEDS = {
    "heic": ("heif",), "heif": ("heif",),
    "avif": ("avif",),
    "jp2": ("jp2",),
    "pdf": ("mupdf",), "svg": ("mupdf",),
    "epub": ("mupdf",), "mobi": ("mupdf",), "fb2": ("mupdf",),
    "cbz": ("mupdf",), "xps": ("mupdf",),
    "docx": ("mammoth",),
    "md": ("markdown",),
    "html": ("mupdf",),
    "xlsx": ("openpyxl",),
    "yaml": ("yaml",),
    "xml": ("xml",),
}

TARGET_NEEDS = {
    (IMAGE, "pdf"): (),

    (DOCUMENT, "docx"): ("pydocx", "mupdf"),
}

ENCODERS = {
    "mp4": "libx264", "mkv": "libx264", "mov": "libx264", "avi": "libx264",
    "webm": "libvpx-vp9",
    "gif": "gif", "apng": "apng", "webp": "libwebp",
    "mp3": "libmp3lame", "ogg": "libvorbis", "opus": "libopus",
    "m4a": "aac", "aac": "aac", "flac": "flac", "wav": "pcm_s16le",
    "aiff": "pcm_s16be", "wma": "wmav2", "ac3": "ac3",
    "srt": "srt", "vtt": "webvtt", "ass": "ass",
}

FFMPEG_CATEGORIES = (VIDEO, AUDIO, SUBTITLE)

_MET: set[str] = set()
_ENCODERS_PRESENT: set[str] = set()
_PROBED = False

def _module(name: str) -> bool:
    try:
        return importlib.util.find_spec(name) is not None
    except (ImportError, ValueError):
        return False

def _pillow_feature(name: str) -> bool:
    try:
        from PIL import features
        return bool(features.check(name))
    except Exception:
        return False

def _ffmpeg_encoders() -> set[str]:
    """The encoder names this ffmpeg was built with."""
    if not shutil.which("ffmpeg"):
        return set()
    try:
        proc = subprocess.run(
            ["ffmpeg", "-hide_banner", "-encoders"],
            capture_output=True, text=True, timeout=20,
        )
    except Exception:
        logger.warning("Could not ask ffmpeg what it can encode; ffmpeg targets are off.")
        return set()
    names = set()
    for line in proc.stdout.splitlines():
        parts = line.split()

        if len(parts) >= 2 and len(parts[0]) == 6 and parts[0][0] in "VAS":
            names.add(parts[1])
    return names

def probe() -> set[str]:
    """Work out what this host can do. Idempotent; called once from main()."""
    global _PROBED, _MET, _ENCODERS_PRESENT
    found = set()
    if _module("pillow_heif"):
        found.add("heif")
    if _pillow_feature("avif") or _module("pillow_heif"):
        found.add("avif")
    if _pillow_feature("jpg_2000"):
        found.add("jp2")
    for requirement, module in (
        ("mupdf", "pymupdf"), ("mammoth", "mammoth"), ("markdown", "markdown"),
        ("openpyxl", "openpyxl"), ("yaml", "yaml"), ("xml", "defusedxml"),
        ("pydocx", "docx"),
    ):
        if _module(module):
            found.add(requirement)
    _ENCODERS_PRESENT = _ffmpeg_encoders()
    if _ENCODERS_PRESENT:
        found.add("ffmpeg")
    _MET = found
    _PROBED = True
    logger.info(
        "Converters available: %s; %d ffmpeg encoders.",
        ", ".join(sorted(found)) or "Pillow only", len(_ENCODERS_PRESENT),
    )
    absent = [name for name in ("mupdf", "heif", "openpyxl", "yaml", "markdown", "mammoth",
                                "pydocx")
              if name not in found]
    if absent:
        logger.warning(
            "Not installed, so those formats are not offered: %s. "
            "pip install -r requirements.txt puts them back.", ", ".join(absent),
        )
    return found

def _ensure_probed() -> None:
    if not _PROBED:
        probe()

def met(requirement: str) -> bool:
    _ensure_probed()
    return requirement in _MET

def _needs_are_met(needs) -> bool:
    _ensure_probed()
    return all(need in _MET for need in needs)

def _encoder_ok(category: str, target: str) -> bool:
    """True unless ffmpeg is the one who would have to write this and cannot."""
    if category not in FFMPEG_CATEGORIES:
        return True
    encoder = ENCODERS.get(target)
    if encoder is None:
        return True
    return encoder in _ENCODERS_PRESENT

def normalise(ext: str | None) -> str:
    """`.JPEG` -> `jpg`. The only way an extension should ever enter."""
    if not ext:
        return ""
    ext = ext.lower().lstrip(".").strip()
    return ALIASES.get(ext, ext)

def category_of(ext: str, animated: bool | None = None) -> str | None:
    """Which category a source file belongs to, or None if it is not one we read."""
    ext = normalise(ext)
    if ext in ANIMATED_CAPABLE and animated is not None:
        return VIDEO if animated else IMAGE
    if ext == "gif":
        return VIDEO if GIF_IS_VIDEO else IMAGE
    return CATEGORY_OF.get(ext)

def is_supported_source(ext: str) -> bool:
    ext = normalise(ext)
    if category_of(ext) is None:
        return False
    return _needs_are_met(FORMAT_NEEDS.get(ext, ()))

def declared_targets(ext: str, animated: bool | None = None) -> tuple[str, ...]:
    """Everything this format converts to on a host that has everything."""
    ext = normalise(ext)
    if ext in ANIMATED_CAPABLE and animated is not None:
        listed = ANIMATED_TARGETS if animated else TARGETS_BY_CATEGORY[IMAGE]
    else:
        listed = TARGETS_BY_FORMAT.get(ext)
        if listed is None:
            listed = TARGETS_BY_CATEGORY.get(category_of(ext), ())
    return tuple(target for target in listed if target != ext)

def targets_for(ext: str, animated: bool | None = None) -> list[str]:
    """Everything this format converts to *here*."""
    ext = normalise(ext)
    category = category_of(ext, animated)
    if category is None or not is_supported_source(ext):
        return []
    out = []
    for target in declared_targets(ext, animated):
        needs = TARGET_NEEDS.get((category, target), FORMAT_NEEDS.get(target, ()))
        if _needs_are_met(needs) and _encoder_ok(category, target):
            out.append(target)
    return out

def popular_targets(ext: str, animated: bool | None = None) -> list[str]:
    """The first screenful. A short list stays whole -- there is no point in hiding two buttons behind a button."""
    targets = targets_for(ext, animated)
    if len(targets) <= POPULAR_COUNT + 2:
        return targets
    return targets[:POPULAR_COUNT]

def supported_sources(category: str) -> list[str]:
    """Every extension of this category this host can read, in the order it is declared -- which is roughly by how common it is."""
    return [ext for ext in SOURCES.get(category, ())
            if category_of(ext) == category and is_supported_source(ext)]

def supported_targets(category: str) -> list[str]:
    """Every format this host can write, for sources of this category."""
    seen = []
    for source in supported_sources(category):
        for target in targets_for(source):
            if target not in seen:
                seen.append(target)
    declared = [target for target in TARGETS_BY_CATEGORY.get(category, ()) if target in seen]
    return declared + [target for target in seen if target not in declared]

def counts() -> tuple[int, int, int]:
    """(formats in, formats out, source-to-target pairs) -- for /formats, which should not contain a number somebody has to remember to update."""
    sources, targets, pairs = set(), set(), 0
    for category in CATEGORIES:
        for source in supported_sources(category):
            available = targets_for(source)
            if not available:
                continue
            sources.add(source)
            targets.update(available)
            pairs += len(available)
    return len(sources), len(targets), pairs

def find_format(text: str) -> str | None:
    """`png`, `.PNG`, `png file`, `image/png` -> `png`, when we know it."""
    if not text:
        return None
    word = text.strip().lower().rsplit("/", 1)[-1].strip().lstrip(".")
    word = word.split()[0] if word.split() else ""
    word = normalise(word)
    return word if word in CATEGORY_OF else None

MAX_PAGES = int(os.environ.get("CONVERT_MAX_PAGES") or 200)
MAX_ROWS = int(os.environ.get("CONVERT_MAX_ROWS") or 100_000)
MAX_CELLS = int(os.environ.get("CONVERT_MAX_CELLS") or 2_000_000)

MAX_BATCH = int(os.environ.get("CONVERT_MAX_BATCH") or 20)

# ─── module: convert_bot.convert_utils ───────────────────────────────────────
"""The conversions themselves, and the Telegram Stars pricing table."""
from __future__ import annotations

import os
import shutil
import subprocess
import zipfile
from pathlib import Path
from typing import NamedTuple

from PIL import Image

import big_files
import formats

class Result(NamedTuple):
    """One file to send, and enough about it to caption it honestly."""
    path: Path
    items: int = 1
    zipped: bool = False

class ConversionError(Exception):
    pass

PROGRESS = None

def report(done: float, total: float) -> None:
    if PROGRESS is None or not total:
        return
    try:
        PROGRESS(min(done, total), total)
    except Exception:
        pass

_LOCAL_BOT_API = bool(os.environ.get("LOCAL_BOT_API_URL"))
big_files.configure("CBOT", "convertbot")
BIG_FILES = _LOCAL_BOT_API or big_files.CONFIGURED
BOT_API_MAX_MB = big_files.BOT_API_DOWNLOAD_BYTES // (1024 * 1024)
MAX_FILE_MB = int(os.environ.get("CONVERT_MAX_FILE_MB")
                  or (big_files.MTPROTO_MAX_BYTES // (1024 * 1024) if BIG_FILES else BOT_API_MAX_MB))
if not BIG_FILES:
    MAX_FILE_MB = min(MAX_FILE_MB, BOT_API_MAX_MB)

def _big_files_refused() -> None:
    """Telegram refused the API id and hash (big_files.verify): promise, and take, only what the Bot API allows -- as if they had never been set."""
    global BIG_FILES, MAX_FILE_MB
    if _LOCAL_BOT_API:
        return
    BIG_FILES = False
    MAX_FILE_MB = min(MAX_FILE_MB, BOT_API_MAX_MB)

big_files.on_disabled(_big_files_refused)

def max_file_mb_in(chat_id: int) -> int:
    """The largest upload taken in this chat."""
    if _LOCAL_BOT_API or big_files.usable(chat_id):
        return MAX_FILE_MB
    return min(MAX_FILE_MB, BOT_API_MAX_MB)

Image.MAX_IMAGE_PIXELS = 80_000_000

MAX_FRAMES = int(os.environ.get("CONVERT_MAX_FRAMES") or 1200)

PAGE_DPI = int(os.environ.get("CONVERT_PAGE_DPI") or 150)

_MB = 1024 * 1024

FREE_UNDER_MB = 3

def price_for_size(size_bytes: int) -> int | None:
    """None if a file this size cannot be taken at all (over MAX_FILE_MB); otherwise what converting it into something its own size would cost."""
    if size_bytes > MAX_FILE_MB * _MB:
        return None
    return _price(size_bytes, size_bytes, 1.0)

BILLABLE = (os.environ.get("CONVERT_BILLABLE") or "sum").strip().lower()

_RATIO_BY_PAIR = {

    ("docx", "pdf"): 20.0, ("docx", "html"): 1.2, ("docx", "md"): 0.3, ("docx", "txt"): 0.15,
    ("md", "pdf"): 4.0, ("md", "html"): 1.3, ("md", "txt"): 0.9,
    ("txt", "pdf"): 3.0, ("txt", "html"): 1.2,
    ("html", "pdf"): 2.0, ("html", "txt"): 0.3,
    ("pdf", "txt"): 0.05, ("pdf", "html"): 0.3, ("pdf", "svg"): 1.5,
    ("epub", "pdf"): 2.5, ("epub", "txt"): 0.25, ("epub", "html"): 0.7,
    ("mobi", "pdf"): 2.5, ("mobi", "txt"): 0.25, ("mobi", "html"): 0.7,
    ("fb2", "pdf"): 2.5, ("fb2", "txt"): 0.3, ("fb2", "html"): 0.8,
    ("xps", "pdf"): 1.1, ("cbz", "pdf"): 1.05,
    ("pdf", "docx"): 0.6,

    ("wav", "mp3"): 0.10, ("wav", "m4a"): 0.09, ("wav", "ogg"): 0.10,
    ("wav", "opus"): 0.07, ("wav", "aac"): 0.09, ("wav", "flac"): 0.6,
    ("flac", "mp3"): 0.25, ("flac", "wav"): 1.7, ("flac", "m4a"): 0.22,
    ("flac", "opus"): 0.18, ("flac", "ogg"): 0.25, ("flac", "aac"): 0.22,
    ("aiff", "mp3"): 0.10, ("aiff", "flac"): 0.6,
    ("mp3", "wav"): 9.0, ("mp3", "flac"): 5.0, ("mp3", "aiff"): 9.0,
    ("m4a", "wav"): 9.0, ("ogg", "wav"): 9.0, ("opus", "wav"): 12.0,

    ("png", "jpg"): 0.15, ("png", "webp"): 0.25, ("png", "avif"): 0.15,
    ("jpg", "png"): 3.5, ("jpg", "bmp"): 12.0, ("jpg", "tiff"): 10.0,
    ("svg", "png"): 20.0, ("svg", "jpg"): 8.0, ("svg", "pdf"): 1.5,
    ("heic", "jpg"): 2.0, ("heic", "png"): 6.0,
}

_RATIO_BY_CATEGORY = {
    (formats.VIDEO, "gif"): 6.0, (formats.VIDEO, "apng"): 9.0, (formats.VIDEO, "webp"): 2.0,
    (formats.VIDEO, "mp4"): 0.9, (formats.VIDEO, "webm"): 0.7, (formats.VIDEO, "mkv"): 1.02,
    (formats.VIDEO, "mov"): 1.05, (formats.VIDEO, "avi"): 1.3,

    (formats.VIDEO, "mp3"): 0.10, (formats.VIDEO, "m4a"): 0.09, (formats.VIDEO, "aac"): 0.09,
    (formats.VIDEO, "ogg"): 0.10, (formats.VIDEO, "opus"): 0.07,
    (formats.VIDEO, "wav"): 0.9, (formats.VIDEO, "flac"): 0.5,
    (formats.VIDEO, "png"): 0.05, (formats.VIDEO, "jpg"): 0.02,
    (formats.IMAGE, "pdf"): 1.05, (formats.IMAGE, "ico"): 0.05, (formats.IMAGE, "icns"): 0.1,
    (formats.DATA, "xlsx"): 0.6, (formats.DATA, "pdf"): 2.0, (formats.DATA, "html"): 1.4,
    (formats.DATA, "md"): 1.1, (formats.DATA, "json"): 1.3, (formats.DATA, "yaml"): 1.1,
    (formats.SUBTITLE, "txt"): 0.7,
}

_RATIO_BY_TARGET = {
    "png": 3.0, "bmp": 10.0, "ppm": 10.0, "tga": 8.0, "tiff": 5.0, "sgi": 8.0,
    "jpg": 0.4, "webp": 0.4, "avif": 0.3, "heic": 0.4, "jp2": 0.8, "qoi": 2.5,
    "gif": 1.5, "ico": 0.05, "icns": 0.1,
    "txt": 0.4, "md": 0.5, "html": 1.2, "pdf": 2.0, "svg": 1.2,
    "wav": 8.0, "flac": 4.0, "aiff": 8.0, "mp3": 0.8, "m4a": 0.8, "aac": 0.8,
    "ogg": 0.8, "opus": 0.6, "wma": 0.9, "ac3": 1.0,
    "docx": 0.8,
}

_BYTES_PER_PAGE = {"png": 350_000, "jpg": 120_000, "webp": 90_000,
                   "tiff": 1_200_000, "svg": 60_000, "pdf": 90_000}

_WORK_BY_TARGET = {
    "mp4": 2.0, "webm": 2.4, "mkv": 2.0, "mov": 2.0, "avi": 1.6,
    "apng": 1.8, "gif": 1.5, "avif": 1.8, "heic": 1.3,
}
_WORK_BY_CATEGORY = {

    (formats.DOCUMENT, "pdf"): 1.4, (formats.DOCUMENT, "png"): 1.4,
    (formats.DOCUMENT, "jpg"): 1.4, (formats.DOCUMENT, "webp"): 1.4,
    (formats.DOCUMENT, "tiff"): 1.5, (formats.DOCUMENT, "svg"): 1.4,
    (formats.DOCUMENT, "docx"): 1.4,

    (formats.VIDEO, "png"): 1.0, (formats.VIDEO, "jpg"): 1.0,
    (formats.VIDEO, "mp3"): 1.1, (formats.VIDEO, "m4a"): 1.1,
}

def work_factor(src_ext: str, target_ext: str) -> float:
    """How much hardware this pair costs for its size, as a multiplier."""
    src_ext, target_ext = formats.normalise(src_ext), formats.normalise(target_ext)
    category = formats.CATEGORY_OF.get(src_ext)
    if (category, target_ext) in _WORK_BY_CATEGORY:
        return _WORK_BY_CATEGORY[(category, target_ext)]
    return _WORK_BY_TARGET.get(target_ext, 1.0)

def estimate_output_bytes(src_ext: str, target_ext: str, in_bytes: int,
                          pages: "int | None" = None) -> int:
    """About how big the result will be. Never exact, and never presented as exact -- see the note above."""
    src_ext, target_ext = formats.normalise(src_ext), formats.normalise(target_ext)
    if pages and target_ext in _BYTES_PER_PAGE and \
            formats.CATEGORY_OF.get(src_ext) == formats.DOCUMENT:
        return max(1, int(pages * _BYTES_PER_PAGE[target_ext]))
    category = formats.CATEGORY_OF.get(src_ext)
    ratio = _RATIO_BY_PAIR.get((src_ext, target_ext))
    if ratio is None:
        ratio = _RATIO_BY_CATEGORY.get((category, target_ext))
    if ratio is None:
        ratio = _RATIO_BY_TARGET.get(target_ext, 1.0)
    return max(1, int(in_bytes * ratio))

def billable_bytes(in_bytes: int, out_bytes: int) -> int:
    """What the price is charged on. See BILLABLE."""
    return max(in_bytes, out_bytes) if BILLABLE == "max" else in_bytes + out_bytes

def price_for_job(src_ext: str, target_ext: str, in_bytes: int,
                  pages: "int | None" = None) -> tuple:
    """(credits, estimated output bytes) for one conversion."""
    out_bytes = estimate_output_bytes(src_ext, target_ext, in_bytes, pages)
    return _price(in_bytes, out_bytes, work_factor(src_ext, target_ext)), out_bytes

def _price(in_bytes: int, out_bytes: int, work: float) -> int:
    import jobs
    import pricing
    billable = billable_bytes(in_bytes, out_bytes)
    if billable < FREE_UNDER_MB * _MB:
        return 0
    return pricing.price(pricing.cost(jobs.time_limit_for(billable, work, paid=True),
                                      jobs.big_send_limit(out_bytes or in_bytes)))

def document_pages(path, src_ext: str) -> "int | None":
    """How many pages a document has, read from its index rather than by rendering it."""
    if formats.CATEGORY_OF.get(formats.normalise(src_ext)) != formats.DOCUMENT:
        return None
    try:
        import pymupdf
    except ImportError:
        return None
    try:
        with pymupdf.open(path) as doc:
            return int(doc.page_count)
    except Exception:
        return None

def least_paid_price() -> int:
    """What the cheapest paid conversion costs: the shortest time a paid job is given and the least it may send back."""
    import jobs
    import pricing
    return pricing.price(pricing.cost(jobs.TIME_BASE_S, jobs.SEND_LIMIT_BYTES))

def price_tiers_text(lang: str = "en") -> str:
    """The price list: free under FREE_UNDER_MB, the least a paid conversion costs, one worked example, and the most anything can -- every number worked out by the same rule a quote is."""
    import i18n
    import jobs
    import pricing

    example_mb = min(100, MAX_FILE_MB)
    example, _ = price_for_job("mp4", "mp4", example_mb * _MB)
    most = pricing.price(pricing.cost(jobs.TIME_LIMIT_S, jobs.big_send_limit(MAX_FILE_MB * _MB)))
    lines = [
        f"  0–{FREE_UNDER_MB} MB: " + i18n.t(lang, "price_free"),
        "  " + i18n.t(lang, "price_from", price=least_paid_price()),
        "  " + i18n.t(lang, "price_example", mb=example_mb, price=example),
        "  " + i18n.t(lang, "price_at_most", mb=MAX_FILE_MB, price=most),
    ]
    note = "" if BIG_FILES else " " + i18n.t(lang, "price_cloud_api_note")
    lines.append(f"  over {MAX_FILE_MB} MB: " + i18n.t(lang, "price_not_supported") + note)
    lines.append("  " + i18n.t(lang, "price_quoted_first"))
    return "\n".join(lines)

_OPENERS_REGISTERED = False

def register_openers() -> None:
    """Teach Pillow the formats that arrive as a plugin."""
    global _OPENERS_REGISTERED
    if _OPENERS_REGISTERED:
        return
    try:
        import pillow_heif
        pillow_heif.register_heif_opener()
    except Exception:
        pass
    _OPENERS_REGISTERED = True

def _open_image(path) -> Image.Image:
    register_openers()
    try:
        img = Image.open(path)
        img.load()
        return img
    except Exception as exc:
        raise ConversionError(f"Couldn't read that as an image: {exc}") from exc

def gif_is_animated(path) -> bool:
    """Whether a GIF, WEBP or APNG actually has more than one frame."""
    try:
        register_openers()
        with Image.open(path) as img:
            return getattr(img, "n_frames", 1) > 1
    except Exception:
        return False

_IMAGE_TARGETS = {
    "png":  ("RGBA", "PNG", {}),
    "jpg":  ("RGB", "JPEG", {"quality": 92}),
    "webp": ("RGBA", "WEBP", {"quality": 92}),
    "bmp":  ("RGB", "BMP", {}),
    "tiff": ("RGB", "TIFF", {"compression": "tiff_lzw"}),
    "avif": ("RGBA", "AVIF", {"quality": 80}),
    "heic": ("RGB", "HEIF", {"quality": 80}),
    "jp2":  ("RGB", "JPEG2000", {"quality_mode": "rates", "quality_layers": [20]}),
    "tga":  ("RGB", "TGA", {}),
    "ppm":  ("RGB", "PPM", {}),
    "qoi":  ("RGB", "QOI", {}),
    "gif":  ("P", "GIF", {}),

    "ico":  ("RGBA", "ICO", {"sizes": [(256, 256), (128, 128), (64, 64), (32, 32), (16, 16)]}),
    "icns": ("RGBA", "ICNS", {}),
}

def _flatten_onto_white(img: Image.Image) -> Image.Image:
    """A transparent PNG saved as JPEG has to become *something* where it was see-through, and black is the one answer nobody wants."""
    background = Image.new("RGB", img.size, (255, 255, 255))
    rgba = img.convert("RGBA")
    background.paste(rgba, mask=rgba.split()[-1])
    return background

def convert_image(in_path, target_ext: str, out_path) -> None:
    target_ext = formats.normalise(target_ext)
    spec = _IMAGE_TARGETS.get(target_ext)
    if spec is None:
        raise ConversionError(f"Unsupported image target format: {target_ext}")
    mode, pillow_format, kwargs = spec

    img = _open_image(in_path)
    try:
        if mode == "RGB" and img.mode in ("RGBA", "LA", "P"):
            img = _flatten_onto_white(img)
        elif mode == "P":
            img = img.convert("RGB").convert("P", palette=Image.ADAPTIVE)
        else:
            img = img.convert(mode)
        if target_ext == "icns":

            img = img.convert("RGBA").resize((1024, 1024))
        img.save(out_path, format=pillow_format, **kwargs)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Image conversion failed: {exc}") from exc
    finally:
        img.close()

def images_to_pdf(in_paths, out_path) -> int:
    """Every picture given, in order, one to a page."""
    pages: list[Image.Image] = []
    try:
        for path in in_paths:
            img = _open_image(path)
            try:

                if img.mode in ("RGBA", "LA", "P"):
                    pages.append(_flatten_onto_white(img))
                else:
                    pages.append(img.convert("RGB"))
            finally:

                img.close()
        if not pages:
            raise ConversionError("No images to put in the PDF.")
        pages[0].save(
            out_path, format="PDF", save_all=True, append_images=pages[1:],
            resolution=PAGE_DPI,
        )
        return len(pages)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Couldn't build the PDF: {exc}") from exc
    finally:
        for page in pages:
            try:
                page.close()
            except Exception:
                pass

FFMPEG_TIMEOUT_S = int(os.environ.get("FFMPEG_TIMEOUT_SECONDS", "600"))
FFMPEG_THREADS = os.environ.get("FFMPEG_THREADS", "2")
FFMPEG_CAP_IN = ["-threads", FFMPEG_THREADS, "-filter_threads", FFMPEG_THREADS,
                 "-filter_complex_threads", FFMPEG_THREADS]
FFMPEG_CAP_OUT = ["-threads", FFMPEG_THREADS]

_EVEN = "scale=trunc(iw/2)*2:trunc(ih/2)*2"
_H264 = ["-c:v", "libx264", "-preset", "veryfast", "-crf", "23",
         "-pix_fmt", "yuv420p", "-vf", _EVEN, "-c:a", "aac"]

_FFMPEG_ARGS = {
    "mp4":  _H264 + ["-movflags", "+faststart"],
    "mov":  _H264,
    "mkv":  _H264,
    "avi":  ["-c:v", "mpeg4", "-qscale:v", "5", "-vf", _EVEN, "-c:a", "libmp3lame", "-q:a", "4"],

    "webm": ["-c:v", "libvpx-vp9", "-crf", "32", "-b:v", "0", "-row-mt", "1",
             "-pix_fmt", "yuv420p", "-vf", _EVEN, "-c:a", "libopus"],

    "gif":  ["-vf", "fps=15,scale=480:-1:flags=lanczos,split[a][b];[a]palettegen[p];[b][p]paletteuse", "-loop", "0"],
    "webp": ["-c:v", "libwebp", "-vf", "fps=15,scale=480:-1:flags=lanczos", "-loop", "0", "-q:v", "70", "-an"],
    "apng": ["-c:v", "apng", "-vf", "fps=15,scale=480:-1:flags=lanczos", "-plays", "0", "-an"],

    "mp3":  ["-vn", "-c:a", "libmp3lame", "-q:a", "2"],
    "wav":  ["-vn", "-c:a", "pcm_s16le"],
    "aiff": ["-vn", "-c:a", "pcm_s16be"],
    "ogg":  ["-vn", "-c:a", "libvorbis", "-q:a", "5"],
    "opus": ["-vn", "-c:a", "libopus", "-b:a", "128k"],
    "m4a":  ["-vn", "-c:a", "aac", "-b:a", "192k"],
    "aac":  ["-vn", "-c:a", "aac", "-b:a", "192k"],
    "flac": ["-vn", "-c:a", "flac"],
    "wma":  ["-vn", "-c:a", "wmav2", "-b:a", "192k"],
    "ac3":  ["-vn", "-c:a", "ac3", "-b:a", "192k"],
    "srt":  ["-c:s", "srt"],
    "vtt":  ["-c:s", "webvtt"],
    "ass":  ["-c:s", "ass"],

    "png":  ["-frames:v", "1", "-an"],
    "jpg":  ["-frames:v", "1", "-q:v", "2", "-an"],
}

def _ffmpeg_available() -> bool:
    return shutil.which("ffmpeg") is not None

_HDR_TRANSFERS = {"smpte2084", "arib-std-b67"}
_TONEMAP = ("zscale=t=linear:npl=100,format=gbrpf32le,zscale=p=bt709,"
            "tonemap=tonemap=hable:desat=0,zscale=t=bt709:m=bt709:r=tv,format=yuv420p")
_SDR_VIDEO = {"mp4", "mov", "mkv", "avi", "webm"}
_SDR_TARGETS = _SDR_VIDEO | {"gif", "webp", "apng", "png", "jpg"}
_ANIMATED_SCALE = "fps=15,scale=480:-1:flags=lanczos"
_filters_seen: "dict[str, bool]" = {}

def _ffmpeg_has_filters(*names: str) -> bool:
    """Whether this ffmpeg was built with every one of these filters. Asked once per process."""
    if not _filters_seen:
        try:
            listing = subprocess.run(["ffmpeg", "-hide_banner", "-filters"], capture_output=True,
                                     text=True, timeout=20).stdout
        except Exception:
            listing = ""
        for line in listing.splitlines():
            parts = line.split()
            if len(parts) >= 2:
                _filters_seen[parts[1]] = True
        _filters_seen.setdefault("", False)
    return all(_filters_seen.get(name) for name in names)

def _is_hdr(in_path) -> bool:
    """Whether the first video stream is on an HDR curve, from ffprobe."""
    if shutil.which("ffprobe") is None:
        return False
    try:
        probe = subprocess.run(
            ["ffprobe", "-v", "error", "-select_streams", "v:0", "-show_entries", "stream=color_transfer",
             "-of", "default=noprint_wrappers=1:nokey=1", str(in_path)],
            capture_output=True, text=True, timeout=20)
    except Exception:
        return False
    return (probe.stdout or "").strip().lower() in _HDR_TRANSFERS

def _tonemapped(args: list, target_ext: str) -> list:
    """The same arguments, with the tonemap in the video filter -- after the downscale for the small animated targets, where it costs a fraction of what it would at 4K, and first everywhere else -- and the result labelled as the BT.709 it now is."""
    args = list(args)
    if "-vf" in args:
        at = args.index("-vf") + 1
        chain = args[at]
        if chain.startswith(_ANIMATED_SCALE):
            args[at] = _ANIMATED_SCALE + "," + _TONEMAP + chain[len(_ANIMATED_SCALE):]
        else:
            args[at] = _TONEMAP + "," + chain
    else:
        args = ["-vf", _TONEMAP] + args
    if target_ext in _SDR_VIDEO:
        args += ["-colorspace", "bt709", "-color_primaries", "bt709", "-color_trc", "bt709"]
    return args

def convert_media_ffmpeg(in_path, target_ext: str, out_path) -> None:
    target_ext = formats.normalise(target_ext)
    args = _FFMPEG_ARGS.get(target_ext)
    if args is None:
        raise ConversionError(f"Unsupported target format: {target_ext}")
    if not _ffmpeg_available():
        raise ConversionError(
            "ffmpeg isn't installed on this host, so video, audio and subtitle "
            "conversion isn't available."
        )
    if not os.path.getsize(in_path):
        raise ConversionError("That file came through empty -- try sending it again.")
    if target_ext in _SDR_TARGETS and _is_hdr(in_path) and _ffmpeg_has_filters("zscale", "tonemap"):
        args = _tonemapped(args, target_ext)

    duration = _media_duration(in_path) if PROGRESS is not None else None
    if duration:
        _ffmpeg_with_progress(in_path, args, out_path, duration)
        return

    cmd = ["ffmpeg", "-nostdin", "-y", *FFMPEG_CAP_IN, "-i", str(in_path)]
    cmd += args + FFMPEG_CAP_OUT
    cmd.append(str(out_path))

    try:
        proc = subprocess.run(cmd, capture_output=True, text=True, timeout=FFMPEG_TIMEOUT_S)
    except subprocess.TimeoutExpired:
        raise ConversionError(
            f"Conversion took longer than {FFMPEG_TIMEOUT_S}s and was stopped. "
            "Try a shorter or smaller file."
        ) from None
    if proc.returncode != 0 or not os.path.exists(out_path):
        note = (proc.stderr or "").strip()[-500:] or "ffmpeg failed with no output"
        raise ConversionError(f"Conversion failed: {note}")

def _media_duration(in_path) -> "float | None":
    """How long the file plays, in seconds, from ffprobe -- or None, and the progress bar simply does not appear."""
    if shutil.which("ffprobe") is None:
        return None
    try:
        probe = subprocess.run(
            ["ffprobe", "-v", "error", "-show_entries", "format=duration",
             "-of", "default=noprint_wrappers=1:nokey=1", str(in_path)],
            capture_output=True, text=True, timeout=20)
        seconds = float((probe.stdout or "").strip().splitlines()[0])
    except Exception:
        return None
    return seconds if seconds > 0 else None

def _ffmpeg_with_progress(in_path, args, out_path, duration: float) -> None:
    """The same ffmpeg run, reading its -progress output as it goes."""
    import tempfile
    import time
    cmd = ["ffmpeg", "-nostdin", "-y", *FFMPEG_CAP_IN, "-progress", "pipe:1", "-nostats",
           "-i", str(in_path)]
    cmd += args + FFMPEG_CAP_OUT
    cmd.append(str(out_path))
    deadline = time.monotonic() + FFMPEG_TIMEOUT_S
    with tempfile.TemporaryFile() as errors:
        proc = subprocess.Popen(cmd, stdout=subprocess.PIPE, stderr=errors, text=True)
        try:
            for line in proc.stdout:
                key, _, value = line.strip().partition("=")
                if key in ("out_time_us", "out_time_ms") and value.isdigit():
                    report(int(value) / 1_000_000, duration)
                if time.monotonic() > deadline:
                    proc.kill()
                    proc.wait()
                    raise ConversionError(
                        f"Conversion took longer than {FFMPEG_TIMEOUT_S}s and was stopped. "
                        "Try a shorter or smaller file.")
            proc.wait(timeout=max(1, deadline - time.monotonic()))
        except subprocess.TimeoutExpired:
            proc.kill()
            proc.wait()
            raise ConversionError(
                f"Conversion took longer than {FFMPEG_TIMEOUT_S}s and was stopped. "
                "Try a shorter or smaller file.") from None
        if proc.returncode != 0 or not os.path.exists(out_path):
            errors.seek(0)
            note = errors.read().decode("utf-8", "replace").strip()[-500:] or "ffmpeg failed with no output"
            raise ConversionError(f"Conversion failed: {note}")

def _reject_overlong_animation(in_path, src_ext: str) -> None:
    """An animation with more frames than MAX_FRAMES is refused up front."""
    if src_ext not in ("gif", "apng", "webp"):
        return
    try:
        register_openers()
        with Image.open(in_path) as img:
            frames = getattr(img, "n_frames", 1)
    except Exception:
        return
    if frames > MAX_FRAMES:
        raise ConversionError(
            f"That animation has {frames} frames, over the {MAX_FRAMES}-frame "
            "limit. Try a shorter clip."
        )

def zip_files(paths, out_path) -> None:
    """One page per entry, deflated."""
    try:
        with zipfile.ZipFile(out_path, "w", zipfile.ZIP_DEFLATED, compresslevel=6) as archive:
            for path in paths:
                archive.write(path, arcname=Path(path).name)
    except Exception as exc:
        raise ConversionError(f"Couldn't package the results: {exc}") from exc

def _out(work_dir, stem: str, ext: str) -> Path:
    return Path(work_dir) / f"{stem}.{ext}"

def convert(in_paths, src_ext: str, target_ext: str, work_dir, stem: str = "converted") -> Result:
    """Everything the bot converts, dispatched by what the source is."""
    in_paths = [str(path) for path in in_paths]
    if not in_paths:
        raise ConversionError("Nothing to convert.")
    src_ext = formats.normalise(src_ext)
    target_ext = formats.normalise(target_ext)
    category = formats.category_of(src_ext)
    work_dir = Path(work_dir)

    if len(in_paths) > 1:
        return _convert_batch(in_paths, src_ext, target_ext, work_dir, stem)

    in_path = in_paths[0]
    out_path = _out(work_dir, stem, target_ext)

    if src_ext in formats.ANIMATED_CAPABLE:
        category = formats.category_of(src_ext, gif_is_animated(in_path))

    if category == formats.IMAGE and src_ext != "svg":
        if target_ext == "pdf":
            images_to_pdf([in_path], out_path)
        elif target_ext in _IMAGE_TARGETS:
            convert_image(in_path, target_ext, out_path)
        else:

            convert_media_ffmpeg(in_path, target_ext, out_path)
        return Result(out_path)

    if category in (formats.VIDEO, formats.AUDIO, formats.SUBTITLE):
        _reject_overlong_animation(in_path, src_ext)
        if category == formats.SUBTITLE and target_ext == "txt":
            import data_convert
            data_convert.subtitle_to_text(in_path, src_ext, out_path)
            return Result(out_path)
        convert_media_ffmpeg(in_path, target_ext, out_path)
        return Result(out_path)

    if category == formats.DOCUMENT or src_ext == "svg":
        import doc_convert
        return doc_convert.convert(in_path, src_ext, target_ext, work_dir, stem)

    if category == formats.DATA:
        import data_convert
        return data_convert.convert(in_path, src_ext, target_ext, work_dir, stem)

    raise ConversionError(f"Nothing here converts a .{src_ext} file.")

def _convert_batch(in_paths, src_ext, target_ext, work_dir, stem) -> Result:
    """Several files at once: one PDF, or one zip of individual results."""
    if len(in_paths) > formats.MAX_BATCH:
        raise ConversionError(
            f"That is {len(in_paths)} files, over the {formats.MAX_BATCH}-file "
            "limit for one conversion."
        )
    if target_ext == "pdf" and formats.category_of(src_ext) == formats.IMAGE:
        out_path = _out(work_dir, stem, "pdf")
        pages = images_to_pdf(in_paths, out_path)
        return Result(out_path, items=pages)

    produced = []
    for index, path in enumerate(in_paths, start=1):
        report(index - 1, len(in_paths))
        one = convert([path], src_ext, target_ext, work_dir, f"{stem}-{index:02d}")
        produced.append(one.path)
    report(len(in_paths), len(in_paths))
    out_path = _out(work_dir, stem, "zip")
    zip_files(produced, out_path)
    for path in produced:
        Path(path).unlink(missing_ok=True)
    return Result(out_path, items=len(produced), zipped=True)

# ─── module: convert_bot.data_convert ────────────────────────────────────────
"""Structured data: CSV, TSV, JSON, YAML, XLSX and XML -- and the plain-text form of a subtitle file, which has nowhere better to live."""
from __future__ import annotations

import csv
import html as html_module
import io
import json
import re
from pathlib import Path
from typing import NamedTuple

import formats
from convert_utils import Result, ConversionError, zip_files

class Table(NamedTuple):
    headers: list[str]
    rows: list[list]

WRITES = frozenset({"json", "yaml", "xlsx", "csv", "tsv", "md", "html", "pdf", "txt"})

_DELIMITERS = {"csv": ",", "tsv": "\t"}

def _read_text(in_path) -> str:
    try:
        return Path(in_path).read_text(encoding="utf-8-sig", errors="replace")
    except OSError as exc:
        raise ConversionError(f"Couldn't read that file: {exc}") from exc

def _check_size(rows: int, columns: int) -> None:
    if rows > formats.MAX_ROWS:
        raise ConversionError(
            f"That is {rows:,} rows, over the {formats.MAX_ROWS:,}-row limit "
            "for one conversion."
        )
    if rows * max(columns, 1) > formats.MAX_CELLS:
        raise ConversionError(
            f"That is about {rows * columns:,} cells, over the "
            f"{formats.MAX_CELLS:,}-cell limit for one conversion."
        )

def _read_delimited(in_path, src_ext: str) -> Table:
    text = _read_text(in_path)
    reader = csv.reader(io.StringIO(text), delimiter=_DELIMITERS[src_ext])
    rows = [row for row in reader]
    if not rows:
        raise ConversionError("That file is empty.")
    headers = [str(cell) for cell in rows[0]]
    _check_size(len(rows) - 1, len(headers))
    return Table(headers, rows[1:])

def _read_xlsx(in_path) -> list[tuple[str, Table]]:
    """Every sheet in the workbook, named."""
    try:
        import openpyxl
    except ImportError as exc:
        raise ConversionError("Spreadsheets aren't supported on this host.") from exc
    try:

        book = openpyxl.load_workbook(in_path, read_only=True, data_only=True)
    except Exception as exc:
        raise ConversionError(f"Couldn't read that spreadsheet: {exc}") from exc
    sheets = []
    try:
        for sheet in book.worksheets:
            rows = [list(row) for row in sheet.iter_rows(values_only=True)]
            rows = [row for row in rows if any(cell is not None for cell in row)]
            if not rows:
                continue
            headers = ["" if cell is None else str(cell) for cell in rows[0]]
            _check_size(len(rows) - 1, len(headers))
            sheets.append((sheet.title, Table(headers, rows[1:])))
    finally:
        book.close()
    if not sheets:
        raise ConversionError("That spreadsheet has no data in it.")
    return sheets

def _read_json(in_path):
    try:
        return json.loads(_read_text(in_path))
    except json.JSONDecodeError as exc:
        raise ConversionError(f"That isn't valid JSON: line {exc.lineno}, {exc.msg}.") from exc

def _read_yaml(in_path):
    try:
        import yaml
    except ImportError as exc:
        raise ConversionError("YAML isn't supported on this host.") from exc
    try:

        return yaml.safe_load(_read_text(in_path))
    except Exception as exc:
        raise ConversionError(f"That isn't valid YAML: {exc}") from exc

def _element_to_object(element):
    """An XML element as dicts and lists, attributes prefixed with `@`."""
    node = {}
    for name, value in element.attrib.items():
        node[f"@{name}"] = value
    for child in element:
        value = _element_to_object(child)
        if child.tag in node:
            if not isinstance(node[child.tag], list):
                node[child.tag] = [node[child.tag]]
            node[child.tag].append(value)
        else:
            node[child.tag] = value
    text = (element.text or "").strip()
    if text:
        if node:
            node["#text"] = text
        else:
            return text
    return node or None

def _read_xml(in_path):
    try:
        from defusedxml.ElementTree import parse
    except ImportError as exc:
        raise ConversionError("XML isn't supported on this host.") from exc
    try:
        root = parse(str(in_path)).getroot()
    except Exception as exc:
        raise ConversionError(f"That isn't XML this can read: {exc}") from exc
    return {root.tag: _element_to_object(root)}

def _load(in_path, src_ext: str):
    """(object, sheets) -- one of the two is always None."""
    if src_ext in _DELIMITERS:
        return None, [("", _read_delimited(in_path, src_ext))]
    if src_ext == "xlsx":
        return None, _read_xlsx(in_path)
    if src_ext == "json":
        return _read_json(in_path), None
    if src_ext == "yaml":
        return _read_yaml(in_path), None
    if src_ext == "xml":
        return _read_xml(in_path), None
    raise ConversionError(f"Nothing here reads a .{src_ext} file.")

def _cell(value):
    """What a nested value looks like in a cell that can only hold a string."""
    if value is None:
        return ""
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False)
    return value

def as_table(obj) -> Table:
    """A nested object as one table, or a refusal that says why not."""
    if isinstance(obj, list) and obj and all(isinstance(item, dict) for item in obj):
        headers: list[str] = []
        for item in obj:
            for key in item:
                if key not in headers:
                    headers.append(key)
        _check_size(len(obj), len(headers))
        return Table(headers, [[_cell(item.get(key)) for key in headers] for item in obj])

    if isinstance(obj, list) and obj and all(isinstance(item, (list, tuple)) for item in obj):
        width = max(len(row) for row in obj)
        _check_size(len(obj), width)
        headers = [f"column {index + 1}" for index in range(width)]
        return Table(headers, [[_cell(cell) for cell in row] + [""] * (width - len(row))
                               for row in obj])

    if isinstance(obj, dict):

        lists = [value for value in obj.values()
                 if isinstance(value, list) and value and all(isinstance(item, dict) for item in value)]
        if len(lists) == 1:
            return as_table(lists[0])
        _check_size(len(obj), 2)
        return Table(["key", "value"], [[key, _cell(value)] for key, value in obj.items()])

    raise ConversionError(
        "That file isn't shaped like a table -- it has no list of records in "
        "it to make rows out of. JSON, YAML and XML will still convert to "
        "each other, or to PDF."
    )

def _table_to_object(table: Table) -> list[dict]:
    return [dict(zip(table.headers, row)) for row in table.rows]

def _object_of(obj, sheets) -> object:
    """Whatever was loaded, as a nested object."""
    if obj is not None:
        return obj
    if len(sheets) == 1:
        return _table_to_object(sheets[0][1])
    return {name or f"sheet {index + 1}": _table_to_object(table)
            for index, (name, table) in enumerate(sheets)}

def _tables_of(obj, sheets) -> list[tuple[str, Table]]:
    """Whatever was loaded, as one or more tables."""
    if sheets is not None:
        return sheets
    return [("", as_table(obj))]

def _write_delimited(table: Table, out_path, target_ext: str) -> None:
    with open(out_path, "w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, delimiter=_DELIMITERS[target_ext])
        writer.writerow(table.headers)
        writer.writerows(table.rows)

def _write_xlsx(sheets: list[tuple[str, Table]], out_path) -> None:
    try:
        import openpyxl
    except ImportError as exc:
        raise ConversionError("Spreadsheets aren't supported on this host.") from exc
    book = openpyxl.Workbook()
    book.remove(book.active)
    for index, (name, table) in enumerate(sheets):

        title = re.sub(r"[:\\/?*\[\]]", "-", name or f"Sheet{index + 1}")[:31]
        sheet = book.create_sheet(title=title)
        sheet.append(table.headers)
        for row in table.rows:
            sheet.append([_cell(value) for value in row])
    book.save(out_path)

def _markdown_table(table: Table) -> str:
    def escape(value):
        return str("" if value is None else value).replace("|", "\\|").replace("\n", " ")

    lines = ["| " + " | ".join(escape(head) for head in table.headers) + " |",
             "| " + " | ".join("---" for _ in table.headers) + " |"]
    for row in table.rows:
        cells = [escape(value) for value in row]
        cells += [""] * (len(table.headers) - len(cells))
        lines.append("| " + " | ".join(cells[:len(table.headers)]) + " |")
    return "\n".join(lines)

def _html_table(table: Table) -> str:
    def cell(value, tag):
        return f"<{tag}>{html_module.escape(str('' if value is None else value))}</{tag}>"

    head = "<tr>" + "".join(cell(head, "th") for head in table.headers) + "</tr>"
    body = "".join(
        "<tr>" + "".join(cell(value, "td") for value in row) + "</tr>" for row in table.rows
    )
    return f"<table border=\"1\" cellspacing=\"0\" cellpadding=\"4\">\n{head}\n{body}\n</table>"

def _write_yaml(obj, out_path) -> None:
    try:
        import yaml
    except ImportError as exc:
        raise ConversionError("YAML isn't supported on this host.") from exc
    with open(out_path, "w", encoding="utf-8") as handle:
        yaml.safe_dump(obj, handle, allow_unicode=True, sort_keys=False, default_flow_style=False)

def convert(in_path, src_ext: str, target_ext: str, work_dir, stem: str = "converted") -> Result:
    src_ext = formats.normalise(src_ext)
    target_ext = formats.normalise(target_ext)
    work_dir = Path(work_dir)
    out_path = work_dir / f"{stem}.{target_ext}"
    obj, sheets = _load(in_path, src_ext)

    if target_ext == "json":
        out_path.write_text(
            json.dumps(_object_of(obj, sheets), ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        return Result(out_path)

    if target_ext == "yaml":
        _write_yaml(json.loads(json.dumps(_object_of(obj, sheets), default=str)), out_path)
        return Result(out_path)

    if target_ext == "xlsx":
        _write_xlsx(_tables_of(obj, sheets), out_path)
        return Result(out_path)

    if target_ext in _DELIMITERS:
        tables = _tables_of(obj, sheets)
        if len(tables) == 1:
            _write_delimited(tables[0][1], out_path, target_ext)
            return Result(out_path)

        produced = []
        for index, (name, table) in enumerate(tables, start=1):
            safe = re.sub(r"[^\w.-]", "_", name) or f"sheet{index}"
            path = work_dir / f"{safe}.{target_ext}"
            _write_delimited(table, path, target_ext)
            produced.append(path)
        archive = work_dir / f"{stem}.zip"
        zip_files(produced, archive)
        for path in produced:
            path.unlink(missing_ok=True)
        return Result(archive, items=len(produced), zipped=True)

    if target_ext in ("md", "html", "pdf"):
        tables = _tables_of(obj, sheets)
        if target_ext == "md":
            parts = []
            for name, table in tables:
                parts.append(f"## {name}\n" if name and len(tables) > 1 else "")
                parts.append(_markdown_table(table))
            out_path.write_text("\n".join(part for part in parts if part), encoding="utf-8")
            return Result(out_path)
        body = "\n".join(
            (f"<h2>{html_module.escape(name)}</h2>" if name and len(tables) > 1 else "")
            + _html_table(table)
            for name, table in tables
        )
        if target_ext == "html":
            out_path.write_text("<!doctype html>\n<meta charset=\"utf-8\">\n" + body, encoding="utf-8")
            return Result(out_path)
        import doc_convert
        doc_convert.html_to_pdf(body, out_path)
        return Result(out_path)

    if target_ext == "txt":
        out_path.write_text(
            json.dumps(_object_of(obj, sheets), ensure_ascii=False, indent=2, default=str),
            encoding="utf-8",
        )
        return Result(out_path)

    raise ConversionError(f"Nothing here converts a .{src_ext} to .{target_ext}.")

_TAG_RE = re.compile(r"<[^>]+>|\{[^}]*\}")
_TIMING_RE = re.compile(r"^\s*[\d:.,]+\s*-->\s*[\d:.,]+")
_INDEX_RE = re.compile(r"^\s*\d+\s*$")
_ASS_EVENT_RE = re.compile(r"^Dialogue:\s*(?:[^,]*,){9}(.*)$")

def subtitle_to_text(in_path, src_ext: str, out_path) -> None:
    source = _read_text(in_path)
    lines: list[str] = []
    if formats.normalise(src_ext) == "ass":
        for line in source.splitlines():
            match = _ASS_EVENT_RE.match(line.strip())
            if match:

                lines.append(_TAG_RE.sub("", match.group(1)).replace("\\N", " ").strip())
    else:
        for line in source.splitlines():
            stripped = line.strip()
            if (not stripped or _INDEX_RE.match(stripped) or _TIMING_RE.match(stripped)
                    or stripped.upper().startswith(("WEBVTT", "NOTE ", "STYLE"))):
                continue
            lines.append(_TAG_RE.sub("", stripped).strip())
    text = "\n".join(line for line in lines if line)
    if not text:
        raise ConversionError("There were no subtitle lines in that file.")
    Path(out_path).write_text(text, encoding="utf-8")

# ─── module: convert_bot.doc_convert ─────────────────────────────────────────
"""Documents: PDF, EPUB, MOBI, FB2, CBZ, XPS, DOCX, Markdown, HTML, plain text -- and SVG, which is a document in every way that matters here."""
from __future__ import annotations

import html as html_module
import io
from pathlib import Path

import formats
from convert_utils import Result, ConversionError, PAGE_DPI, convert_image, report, zip_files

_MUPDF_NATIVE = ("pdf", "epub", "mobi", "fb2", "cbz", "xps", "svg", "html", "txt")

WRITES = frozenset({"pdf", "png", "jpg", "webp", "tiff", "svg", "txt", "html", "md", "docx"})

_PAGE = "a4"
_MARGIN = 40

def _fitz():
    try:
        import pymupdf
        return pymupdf
    except ImportError as exc:
        raise ConversionError(
            "The document engine isn't installed on this host, so PDF and "
            "e-book conversion isn't available."
        ) from exc

def _open(in_path, src_ext: str):
    """The source as a PyMuPDF document, whatever it started as."""
    pymupdf = _fitz()
    if src_ext in _MUPDF_NATIVE:
        try:
            return pymupdf.open(in_path)
        except Exception as exc:
            raise ConversionError(f"Couldn't read that {src_ext.upper()}: {exc}") from exc

    pdf_bytes = _to_pdf_bytes(in_path, src_ext)
    return pymupdf.open(stream=pdf_bytes, filetype="pdf")

def _page_count(doc) -> int:
    count = doc.page_count
    if count > formats.MAX_PAGES:
        raise ConversionError(
            f"That document is {count} pages, over the {formats.MAX_PAGES}-page "
            "limit for one conversion. Split it and send the part you need."
        )
    return count

def _docx_to_html(in_path) -> str:
    try:
        import mammoth
    except ImportError as exc:
        raise ConversionError("Word documents aren't supported on this host.") from exc
    try:
        with open(in_path, "rb") as handle:
            return mammoth.convert_to_html(handle).value
    except Exception as exc:
        raise ConversionError(f"Couldn't read that .docx: {exc}") from exc

def _md_to_html(in_path) -> str:
    try:
        import markdown
    except ImportError as exc:
        raise ConversionError("Markdown isn't supported on this host.") from exc
    try:
        text = Path(in_path).read_text(encoding="utf-8", errors="replace")
    except OSError as exc:
        raise ConversionError(f"Couldn't read that file: {exc}") from exc
    return markdown.markdown(text, extensions=["tables", "fenced_code", "sane_lists"])

def _source_html(in_path, src_ext: str) -> str:
    if src_ext == "docx":
        return _docx_to_html(in_path)
    if src_ext == "md":
        return _md_to_html(in_path)
    if src_ext == "html":
        return Path(in_path).read_text(encoding="utf-8", errors="replace")
    if src_ext == "txt":
        body = html_module.escape(Path(in_path).read_text(encoding="utf-8", errors="replace"))
        return f"<pre>{body}</pre>"

    doc = _open(in_path, src_ext)
    try:
        pages = _page_count(doc)
        return "\n".join(doc[index].get_text("html") for index in range(pages))
    finally:
        doc.close()

def html_to_pdf(source_html: str, out_path) -> None:
    """Lay HTML out onto pages and write a PDF."""
    pymupdf = _fitz()
    try:
        story = pymupdf.Story(html=f"<html><body>{source_html}</body></html>")
        page_rect = pymupdf.paper_rect(_PAGE)
        frame = page_rect + (_MARGIN, _MARGIN, -_MARGIN, -_MARGIN)

        writer = pymupdf.DocumentWriter(out_path if hasattr(out_path, "write") else str(out_path))
        more = 1
        pages = 0
        while more:
            device = writer.begin_page(page_rect)
            more, _ = story.place(frame)
            story.draw(device)
            writer.end_page()
            pages += 1
            if pages > formats.MAX_PAGES:
                writer.close()
                raise ConversionError(
                    f"That comes to more than {formats.MAX_PAGES} pages, which is "
                    "over the limit for one conversion."
                )
        writer.close()
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Couldn't lay that out as a PDF: {exc}") from exc

def _to_pdf_bytes(in_path, src_ext: str) -> bytes:
    """A PDF of the source, in memory."""
    _fitz()
    if src_ext in _MUPDF_NATIVE:
        doc = _open(in_path, src_ext)
        try:
            _page_count(doc)
            if src_ext == "pdf":
                return doc.tobytes()
            return doc.convert_to_pdf()
        finally:
            doc.close()
    buffer = io.BytesIO()
    html_to_pdf(_source_html(in_path, src_ext), buffer)
    return buffer.getvalue()

def _to_pdf(in_path, src_ext: str, out_path) -> None:
    if src_ext in ("docx", "md"):
        html_to_pdf(_source_html(in_path, src_ext), out_path)
        return
    doc = _open(in_path, src_ext)
    try:
        _page_count(doc)
        Path(out_path).write_bytes(doc.convert_to_pdf())
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Couldn't write that as a PDF: {exc}") from exc
    finally:
        doc.close()

def _render_pages(in_path, src_ext: str, target_ext: str, work_dir, stem: str) -> list[Path]:
    """One picture per page, at PAGE_DPI."""
    doc = _open(in_path, src_ext)
    produced: list[Path] = []
    try:
        pages = _page_count(doc)
        for index in range(pages):
            report(index, pages)
            pixmap = doc[index].get_pixmap(dpi=PAGE_DPI)
            png_path = Path(work_dir) / f"{stem}-page{index + 1:03d}.png"
            pixmap.save(str(png_path))
            if target_ext == "png":
                produced.append(png_path)
                continue
            out_path = png_path.with_suffix(f".{target_ext}")
            try:
                convert_image(png_path, target_ext, out_path)
            finally:
                png_path.unlink(missing_ok=True)
            produced.append(out_path)
    except ConversionError:
        for path in produced:
            path.unlink(missing_ok=True)
        raise
    except Exception as exc:
        for path in produced:
            path.unlink(missing_ok=True)
        raise ConversionError(f"Couldn't render that document: {exc}") from exc
    finally:
        doc.close()
    return produced

def _page_svgs(in_path, src_ext: str, work_dir, stem: str) -> list[Path]:
    """A PDF page as SVG keeps its vectors, which is the whole reason to ask for one rather than a PNG."""
    doc = _open(in_path, src_ext)
    produced: list[Path] = []
    try:
        pages = _page_count(doc)
        for index in range(pages):
            report(index, pages)
            path = Path(work_dir) / f"{stem}-page{index + 1:03d}.svg"
            path.write_text(doc[index].get_svg_image(), encoding="utf-8")
            produced.append(path)
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Couldn't write that page as SVG: {exc}") from exc
    finally:
        doc.close()
    return produced

def _pack(produced: list[Path], target_ext: str, work_dir, stem: str) -> Result:
    """One page comes back as itself; many come back as a zip."""
    if not produced:
        raise ConversionError("That document had no pages in it.")
    if len(produced) == 1:
        final = Path(work_dir) / f"{stem}.{target_ext}"
        produced[0].replace(final)
        return Result(final, items=1)
    archive = Path(work_dir) / f"{stem}.zip"
    zip_files(produced, archive)
    for path in produced:
        path.unlink(missing_ok=True)
    return Result(archive, items=len(produced), zipped=True)

def _to_text(in_path, src_ext: str) -> str:
    if src_ext == "docx":
        try:
            import mammoth
            with open(in_path, "rb") as handle:
                return mammoth.extract_raw_text(handle).value
        except ImportError as exc:
            raise ConversionError("Word documents aren't supported on this host.") from exc
        except Exception as exc:
            raise ConversionError(f"Couldn't read that .docx: {exc}") from exc
    if src_ext == "md":

        doc = _open(in_path, src_ext)
        try:
            return "\n".join(doc[index].get_text() for index in range(_page_count(doc)))
        finally:
            doc.close()
    doc = _open(in_path, src_ext)
    try:
        return "\n".join(doc[index].get_text() for index in range(_page_count(doc)))
    except ConversionError:
        raise
    except Exception as exc:
        raise ConversionError(f"Couldn't extract the text: {exc}") from exc
    finally:
        doc.close()

_HEADING_RATIO = 1.25
_HEADING_2_RATIO = 1.6

_MIN_IMAGE_PIXELS = 64 * 64
_DOCX_PAGE_WIDTH_INCHES = 6.0

def _docx():
    try:
        import docx
        return docx
    except ImportError as exc:
        raise ConversionError(
            "The Word writer isn't installed on this host, so converting to "
            "DOCX isn't available."
        ) from exc

def _body_size(doc) -> float:
    """The size most of this document's text is set in, which is what makes 'larger than the body' a question with an answer."""
    sizes: dict = {}
    for page in doc:
        for block in page.get_text("dict").get("blocks", ()):
            for line in block.get("lines", ()):
                for span in line.get("spans", ()):
                    text = (span.get("text") or "").strip()
                    if not text:
                        continue
                    size = round(float(span.get("size") or 0), 1)
                    sizes[size] = sizes.get(size, 0) + len(text)
    if not sizes:
        return 0.0
    return max(sizes.items(), key=lambda pair: pair[1])[0]

def _line_text(line) -> str:
    return "".join(span.get("text") or "" for span in line.get("spans", ())).strip()

def _to_docx(in_path, src_ext: str, out_path, work_dir) -> None:
    """Everything the PDF actually says, as a Word document. See above for what that does and does not include."""
    docx = _docx()
    pymupdf = _fitz()
    document = docx.Document()
    with _open(in_path, src_ext) as doc:
        _page_count(doc)
        body_size = _body_size(doc)
        wrote_anything = False
        for index, page in enumerate(doc):
            report(index, len(doc))
            if index:
                document.add_page_break()
            for block in sorted(page.get_text("dict").get("blocks", ()),
                                key=lambda b: (round(b.get("bbox", (0, 0, 0, 0))[1]),
                                               round(b.get("bbox", (0, 0, 0, 0))[0]))):
                for line in block.get("lines", ()):
                    text = _line_text(line)
                    if not text:
                        continue
                    spans = [span for span in line.get("spans", ()) if (span.get("text") or "").strip()]
                    size = max((float(span.get("size") or 0) for span in spans), default=body_size)

                    bold = any(int(span.get("flags") or 0) & (1 << 4) for span in spans)
                    if body_size and size >= body_size * _HEADING_2_RATIO:
                        document.add_heading(text, level=1)
                    elif body_size and size >= body_size * _HEADING_RATIO:
                        document.add_heading(text, level=2)
                    else:
                        paragraph = document.add_paragraph()
                        run = paragraph.add_run(text)
                        run.bold = bold
                    wrote_anything = True
            _add_page_images(document, doc, page, pymupdf, work_dir)
    if not wrote_anything:
        raise ConversionError(
            "There is nothing in that file to put into a Word document -- it "
            "looks like scanned pages rather than text. Converting it to PNG "
            "or JPG will give you the pages themselves."
        )
    document.save(str(out_path))

def _add_page_images(document, doc, page, pymupdf, work_dir) -> None:
    """The pictures on one page, in the order the page lists them."""
    import io
    from PIL import Image as PILImage
    try:
        listed = page.get_images(full=True)
    except Exception:
        return
    for number, info in enumerate(listed):
        try:
            data = doc.extract_image(info[0])
            raw, width, height = data["image"], data.get("width", 0), data.get("height", 0)
            if width * height < _MIN_IMAGE_PIXELS:
                continue

            picture = PILImage.open(io.BytesIO(raw))
            if picture.mode not in ("RGB", "L"):
                picture = picture.convert("RGB")
            buffer = io.BytesIO()
            picture.save(buffer, format="PNG")
            buffer.seek(0)
            inches = min(_DOCX_PAGE_WIDTH_INCHES, max(1.0, width / 96))
            from docx.shared import Inches
            document.add_picture(buffer, width=Inches(inches))
        except Exception:
            continue

def convert(in_path, src_ext: str, target_ext: str, work_dir, stem: str = "converted") -> Result:
    src_ext = formats.normalise(src_ext)
    target_ext = formats.normalise(target_ext)
    work_dir = Path(work_dir)
    out_path = work_dir / f"{stem}.{target_ext}"

    if target_ext == "pdf":
        _to_pdf(in_path, src_ext, out_path)
        return Result(out_path)

    if target_ext == "docx":
        _to_docx(in_path, src_ext, out_path, work_dir)
        return Result(out_path)

    if target_ext in ("png", "jpg", "webp", "tiff"):
        return _pack(_render_pages(in_path, src_ext, target_ext, work_dir, stem),
                     target_ext, work_dir, stem)

    if target_ext == "svg":
        return _pack(_page_svgs(in_path, src_ext, work_dir, stem), "svg", work_dir, stem)

    if target_ext == "txt":
        text = _to_text(in_path, src_ext)
        if not text.strip():
            raise ConversionError(
                "There is no text in that file to extract -- it looks like "
                "scanned pages rather than a document. Converting it to PNG "
                "or JPG will give you the pages themselves."
            )
        out_path.write_text(text, encoding="utf-8")
        return Result(out_path)

    if target_ext == "html":
        out_path.write_text(
            "<!doctype html>\n<meta charset=\"utf-8\">\n" + _source_html(in_path, src_ext),
            encoding="utf-8",
        )
        return Result(out_path)

    if target_ext == "md":
        if src_ext != "docx":
            raise ConversionError(f"Nothing here turns a .{src_ext} into Markdown.")
        try:
            import mammoth
            with open(in_path, "rb") as handle:
                out_path.write_text(mammoth.convert_to_markdown(handle).value, encoding="utf-8")
        except ImportError as exc:
            raise ConversionError("Word documents aren't supported on this host.") from exc
        except Exception as exc:
            raise ConversionError(f"Couldn't read that .docx: {exc}") from exc
        return Result(out_path)

    raise ConversionError(f"Nothing here converts a .{src_ext} to .{target_ext}.")

# ─── module: convert_bot.convert_worker ──────────────────────────────────────
"""One conversion, in a process of its own."""
import json
import os
import sys
import traceback
import warnings


PROGRESS_EVERY_S = 1.0

def _emit(ok: bool, **fields) -> None:
    fields["ok"] = ok
    sys.stdout.write("\n" + json.dumps(fields) + "\n")
    sys.stdout.flush()

def _cap_memory(max_memory_mb: int) -> None:
    """Give the worker a MemoryError instead of letting the kernel pick a victim."""
    if not max_memory_mb or os.name != "posix":
        return
    try:
        import resource
        limit = int(max_memory_mb) * 1024 * 1024
        _, hard = resource.getrlimit(resource.RLIMIT_AS)
        resource.setrlimit(resource.RLIMIT_AS, (limit, hard))
    except Exception:
        pass

def main(argv) -> int:
    if len(argv) < 2:
        _emit(False, kind="crash", message="no job was given to the worker")
        return 1
    try:
        job = json.loads(argv[1])
    except ValueError as exc:
        _emit(False, kind="crash", message=f"the job was not valid JSON: {exc}")
        return 1

    _cap_memory(job.get("max_memory_mb") or 0)

    warnings.simplefilter("ignore")

    try:
        import convert_utils
    except Exception as exc:
        traceback.print_exc(file=sys.stderr)
        _emit(False, kind="crash", message=f"could not load the converters: {exc}")
        return 1

    import time
    last = [0.0]

    def progress(done, total):
        now = time.monotonic()
        if now - last[0] >= PROGRESS_EVERY_S or done >= total:
            last[0] = now
            sys.stdout.write(json.dumps({"progress": [round(float(done), 3), round(float(total), 3)]}) + "\n")
            sys.stdout.flush()

    convert_utils.PROGRESS = progress

    try:
        from PIL import Image
        if job.get("max_pixels"):
            Image.MAX_IMAGE_PIXELS = int(job["max_pixels"])
    except Exception:
        pass

    try:
        result = convert_utils.convert(
            job["in_paths"], job["src_ext"], job["target_ext"],
            job["work_dir"], job.get("stem") or "converted",
        )
    except convert_utils.ConversionError as exc:
        _emit(False, kind="refused", message=str(exc))
        return 2
    except MemoryError:
        _emit(False, kind="crash", message="the conversion ran out of memory")
        return 1
    except Exception as exc:
        traceback.print_exc(file=sys.stderr)
        _emit(False, kind="crash", message=f"{type(exc).__name__}: {exc}")
        return 1

    path = str(result.path)
    try:
        size = os.path.getsize(path)
    except OSError as exc:
        _emit(False, kind="crash", message=f"the converter reported a file that is not there: {exc}")
        return 1
    _emit(True, path=path, zipped=bool(getattr(result, "zipped", False)),
          items=int(getattr(result, "items", 1) or 1), bytes=size)
    return 0

if __name__ == "__main__":
    sys.exit(main(sys.argv))

# ─── module: convert_bot.jobs ────────────────────────────────────────────────
"""The conversion queue, and what happens to the credit."""
from __future__ import annotations

import asyncio
import itertools
import json
import logging
import os
import shutil
import signal
import sys
import time
import uuid
from dataclasses import dataclass, field
from pathlib import Path

import big_files

logger = logging.getLogger(__name__)

HERE = Path(__spec__.origin if __spec__ else __file__).resolve().parent
WORKER = HERE / "main.py"
WORKER_ARGS = ("--module", "convert_bot.convert_worker")

LOCAL_BOT_API = bool(os.environ.get("LOCAL_BOT_API_URL"))

MAX_CONCURRENT = max(1, int(os.environ.get("CONVERT_MAX_CONCURRENT") or 1))

TIME_LIMIT_S = int(os.environ.get("CONVERT_TIME_LIMIT_SECONDS") or 600)

FREE_TIME_LIMIT_S = int(os.environ.get("CONVERT_FREE_TIME_LIMIT_SECONDS") or 60)

TIME_BASE_S = int(os.environ.get("CONVERT_TIME_BASE_SECONDS") or 120)
TIME_PER_MB_S = float(os.environ.get("CONVERT_TIME_PER_MB_SECONDS") or 1.0)

def time_limit_for(billable_bytes: int, work: float = 1.0, paid: bool = True) -> int:
    """How long a conversion may run, in seconds."""
    if not paid:
        return FREE_TIME_LIMIT_S
    wanted = TIME_BASE_S + TIME_PER_MB_S * work * billable_bytes / (1024 * 1024)
    return int(min(TIME_LIMIT_S, max(TIME_BASE_S, wanted)))

def job_limit(job: "Job") -> int:
    """The limit this job runs under: its own, or for a job made before it had one, the flat limit of its kind."""
    return job.time_limit or (TIME_LIMIT_S if job.price > 0 else FREE_TIME_LIMIT_S)

PROJECT_AFTER_S = 30
PROJECT_MARGIN = 1.2

PROGRESS_MARKS_S = (120, 300)

MAX_QUEUED_PER_USER = max(1, int(os.environ.get("CONVERT_MAX_QUEUED_PER_USER") or 5))

def _default_user_storage_mb() -> int:
    try:
        import convert_utils
        import formats
        album = formats.MAX_BATCH * convert_utils.MAX_FILE_MB
        return min(album, max(2000, 2 * convert_utils.MAX_FILE_MB)) + 50
    except Exception:
        return 450

USER_STORAGE_MB = int(os.environ.get("CONVERT_USER_STORAGE_MB") or max(120, _default_user_storage_mb()))
STORAGE_MB = int(os.environ.get("CONVERT_STORAGE_MB") or max(1024, USER_STORAGE_MB + 200))
MIN_FREE_MB = int(os.environ.get("CONVERT_MIN_FREE_MB") or 300)

MAX_MEGAPIXELS = int(os.environ.get("CONVERT_MAX_MEGAPIXELS") or 210)

CARD_EVERY_S = 3.0

HOLD_TOUCH_S = 30
HOLD_STALE_S = 120

WORKER_MAX_MEMORY_MB = int(os.environ.get("CONVERT_WORKER_MAX_MEMORY_MB") or 0)

SEND_LIMIT_MB = big_files.BOT_API_SEND_BYTES // (1024 * 1024)
SEND_LIMIT_BYTES = SEND_LIMIT_MB * 1024 * 1024

BIG_SEND_MAX_BYTES = big_files.MTPROTO_MAX_BYTES
SEND_HEADROOM = 3

def big_sends(chat_id: int) -> bool:
    """Whether a result over 50 MB can reach this chat at all."""
    return LOCAL_BOT_API or big_files.usable(chat_id)

def big_send_limit(out_estimate: int) -> int:
    """The largest result a conversion estimated at `out_estimate` bytes may send where big sends can reach, which is what its price covers."""
    return min(BIG_SEND_MAX_BYTES, max(SEND_LIMIT_BYTES, SEND_HEADROOM * max(0, out_estimate or 0)))

def send_limit_for(in_bytes: int, out_estimate: int | None, chat_id: int) -> int:
    """The largest result this conversion may send back, in bytes."""
    if not big_sends(chat_id):
        return SEND_LIMIT_BYTES
    return big_send_limit(out_estimate if out_estimate else in_bytes)

def send_ceiling_mb(chat_id: int) -> int:
    """The most any result could be sent at in this chat, for sentences about a format that is off the menu for being too big."""
    return (BIG_SEND_MAX_BYTES if big_sends(chat_id) else SEND_LIMIT_BYTES) // (1024 * 1024)

LOSSLESS_BYTES_PER_PIXEL = {
    "bmp": 3.0,
    "ppm": 3.0,
    "tga": 3.0,
    "tiff": 0.6,
    "png": 0.5,
    "qoi": 0.5,
}

TARGET_MAX_MEGAPIXELS = {
    "heic": 142.6,
}

USER_CANCELLED = "user_cancelled"
NO_CREDIT = "no_credit"
REFUSED = "refused"
TIMEOUT = "timeout"

TOO_SLOW = "too_slow"
CRASH = "crash"
TOO_LARGE = "too_large"
SEND_FAILED = "send_failed"
INTERRUPTED = "interrupted"

PAUSED = "paused"

REFUNDED = frozenset({REFUSED, TIMEOUT, TOO_SLOW, CRASH, TOO_LARGE, SEND_FAILED, INTERRUPTED})

ALERTED = frozenset({TIMEOUT, CRASH, TOO_LARGE, SEND_FAILED, INTERRUPTED})

_sequence = itertools.count()

@dataclass
class Job:
    user_id: int
    chat_id: int
    paths: list
    src_ext: str
    target_ext: str
    size: int
    price: int
    lang: str
    username: str | None = None
    keep_file: bool = False

    send_limit: int = SEND_LIMIT_BYTES

    time_limit: int = 0
    id: str = field(default_factory=lambda: uuid.uuid4().hex[:10])
    state: str = "queued"
    charged: bool = False
    user_cancelled: bool = False
    task: asyncio.Task | None = None
    proc: object | None = None
    queued_at: float = field(default_factory=time.monotonic)
    seq: int = field(default_factory=lambda: next(_sequence))
    started_at: float | None = None
    extra: dict = field(default_factory=dict)

@dataclass
class Outcome:
    ok: bool
    kind: str | None = None
    message: str = ""
    path: str | None = None
    zipped: bool = False
    items: int = 1
    bytes: int = 0

_jobs: dict[str, Job] = {}
_user_locks: dict[int, asyncio.Lock] = {}
_slots: asyncio.Semaphore | None = None

def reset() -> None:
    """For tests, which run each scenario on a fresh event loop: asyncio primitives belong to the loop they were first used on."""
    global _slots
    _jobs.clear()
    _user_locks.clear()
    _slots = None

def _global_slots() -> asyncio.Semaphore:
    global _slots
    if _slots is None:
        _slots = asyncio.Semaphore(MAX_CONCURRENT)
    return _slots

def get(job_id: str) -> Job | None:
    return _jobs.get(job_id)

def jobs_for(user_id: int) -> list[Job]:
    return sorted((job for job in _jobs.values() if job.user_id == user_id),
                  key=lambda job: job.seq)

def ahead_of(job: Job) -> int:
    """How many jobs will start before this one, across everybody."""
    return sum(1 for other in _jobs.values()
               if other is not job and other.state in ("queued", "running")
               and other.seq < job.seq)

def active_paths() -> set[str]:
    """Every staged file a queued or running job still needs."""
    return {str(path) for job in _jobs.values() for path in job.paths}

def uncharged_total(user_id: int) -> int:
    """Credit this person's queued jobs will need when they start."""
    return sum(job.price for job in _jobs.values()
               if job.user_id == user_id and not job.charged)

def storage_problem(user_id: int, incoming_bytes: int, upload_dir) -> tuple[str, int, int] | None:
    """None if there is room for `incoming_bytes` more; otherwise which limit stops it -- "user", "global" or "disk" -- with the MB used and the limit."""
    upload_dir = Path(upload_dir)
    mb = 1024 * 1024
    used_all = used_user = 0
    prefix = f"{user_id}-"
    if upload_dir.is_dir():
        for path in upload_dir.iterdir():
            try:
                if path.is_file():
                    size = path.stat().st_size
                    used_all += size
                    if path.name.startswith(prefix):
                        used_user += size
            except OSError:
                continue
    if USER_STORAGE_MB and used_user + incoming_bytes > USER_STORAGE_MB * mb:
        return ("user", used_user // mb, USER_STORAGE_MB)
    if STORAGE_MB and used_all + incoming_bytes > STORAGE_MB * mb:
        return ("global", used_all // mb, STORAGE_MB)
    try:
        probe = upload_dir if upload_dir.is_dir() else upload_dir.parent
        free = shutil.disk_usage(probe).free
    except OSError:
        free = None
    if free is not None and MIN_FREE_MB and free - incoming_bytes < MIN_FREE_MB * mb:
        return ("disk", free // mb, MIN_FREE_MB)
    return None

def image_megapixels(path) -> float | None:
    """The size of a picture, read from its header without decoding it."""
    try:
        from PIL import Image
    except Exception:
        return None
    previous = getattr(Image, "MAX_IMAGE_PIXELS", None)
    try:
        Image.MAX_IMAGE_PIXELS = None
        with Image.open(path) as img:
            width, height = img.size
        return (width * height) / 1_000_000
    except Exception:
        return None
    finally:
        try:
            Image.MAX_IMAGE_PIXELS = previous
        except Exception:
            pass

def estimate_output_bytes(megapixels: float | None, target_ext: str) -> int | None:
    ratio = LOSSLESS_BYTES_PER_PIXEL.get(target_ext)
    if megapixels is None or ratio is None:
        return None
    return int(megapixels * 1_000_000 * ratio)

def unsendable_targets(megapixels: float | None, targets, in_bytes: int = 0,
                       chat_id: int = 0) -> list[str]:
    """The formats that would certainly come out too big to send, or that the encoder itself cannot produce at this size."""
    out = []
    for target in targets:
        ceiling = TARGET_MAX_MEGAPIXELS.get(target)
        if megapixels is not None and ceiling is not None and megapixels > ceiling:
            out.append(target)
            continue
        estimate = estimate_output_bytes(megapixels, target)
        if estimate is not None and estimate > send_limit_for(in_bytes, estimate, chat_id):
            out.append(target)
    return out

_OWN_GROUP = {} if sys.platform == "win32" else {"start_new_session": True}

async def _kill(proc) -> None:
    if proc is None or proc.returncode is not None:
        return
    try:
        if sys.platform == "win32":
            killer = await asyncio.create_subprocess_exec(
                "taskkill", "/F", "/T", "/PID", str(proc.pid),
                stdout=asyncio.subprocess.DEVNULL, stderr=asyncio.subprocess.DEVNULL)
            await killer.wait()
        else:
            os.killpg(proc.pid, signal.SIGKILL)
    except ProcessLookupError:
        return
    except Exception:
        logger.debug("Could not kill the conversion worker's group", exc_info=True)
    if proc.returncode is None:
        try:
            proc.kill()
        except ProcessLookupError:
            return
        except Exception:
            logger.debug("Could not kill conversion worker", exc_info=True)
    try:
        await asyncio.wait_for(proc.wait(), timeout=10)
    except Exception:
        logger.debug("Conversion worker did not exit after kill", exc_info=True)

def _parse(returncode, out: bytes, err: bytes) -> Outcome:
    lines = [line for line in (out or b"").decode("utf-8", "replace").splitlines() if line.strip()]
    if lines:
        try:
            data = json.loads(lines[-1])
        except ValueError:
            data = None
        if isinstance(data, dict):
            if data.get("ok"):
                return Outcome(True, None, "", data.get("path"), bool(data.get("zipped")),
                               int(data.get("items") or 1), int(data.get("bytes") or 0))
            return Outcome(False, data.get("kind") or CRASH, str(data.get("message") or ""))
    tail = (err or b"").decode("utf-8", "replace").strip()[-400:]
    if returncode is not None and (returncode < 0 or returncode == 137):
        return Outcome(False, CRASH,
                       f"the worker was killed (exit {returncode}), most likely for memory. {tail}".strip())
    return Outcome(False, CRASH, f"the worker exited {returncode} without a result. {tail}".strip())

async def run_worker(job: Job, work_dir, stem: str, time_limit: int) -> Outcome:
    payload = {
        "in_paths": [str(path) for path in job.paths],
        "src_ext": job.src_ext,
        "target_ext": job.target_ext,
        "work_dir": str(work_dir),
        "stem": stem,
        "max_pixels": MAX_MEGAPIXELS * 1_000_000,
        "max_memory_mb": WORKER_MAX_MEMORY_MB,
    }
    proc = await asyncio.create_subprocess_exec(
        sys.executable, "-u", str(WORKER), *WORKER_ARGS, json.dumps(payload),
        cwd=str(HERE), stdout=asyncio.subprocess.PIPE, stderr=asyncio.subprocess.PIPE, **_OWN_GROUP,
    )
    job.proc = proc
    kept: list[bytes] = []
    errors = bytearray()

    async def follow():

        async def drain():
            while chunk := await proc.stderr.read(65536):
                errors.extend(chunk)
                del errors[:-8000]
        draining = asyncio.ensure_future(drain())
        while True:
            try:
                line = await proc.stdout.readline()
            except ValueError:
                continue
            if not line:
                break
            try:
                data = json.loads(line) if line.startswith(b'{"progress"') else None
            except ValueError:
                data = None
            if data is not None:
                done, total = data["progress"]
                job.extra["progress"] = (done, total)
                elapsed = time.monotonic() - started
                if (elapsed >= PROJECT_AFTER_S and total and done > 0 and "projected_s" not in job.extra
                        and elapsed * total / done > time_limit * PROJECT_MARGIN):

                    job.extra["projected_s"] = elapsed * total / done
                    await _kill(proc)
            else:
                kept.append(line)
        await proc.wait()
        await draining

    started = time.monotonic()
    try:
        await asyncio.wait_for(follow(), timeout=time_limit)
        if "projected_s" in job.extra:
            return Outcome(False, TOO_SLOW, f"on course for {job.extra['projected_s']:.0f}s against a "
                                            f"{time_limit}s limit, stopped after {time.monotonic() - started:.0f}s")
        out, err = b"".join(kept), bytes(errors)
    except asyncio.TimeoutError:
        await _kill(proc)
        return Outcome(False, TIMEOUT, f"ran past the {time_limit}s limit")
    except asyncio.CancelledError:
        await asyncio.shield(_kill(proc))
        raise
    finally:
        job.proc = None
    return _parse(proc.returncode, out, err)

class Runner:
    """What bot.py provides. None of these may raise, except `deliver`, whose exception is the send failure."""

    def work_dir(self) -> Path:
        raise NotImplementedError

    async def refuse_start(self, job: Job) -> bool:
        return False

    async def charge(self, job: Job) -> int | None:
        raise NotImplementedError

    async def started(self, job: Job, balance: int) -> None:
        return None

    async def progress(self, job: Job, stage: int, stages: int, elapsed_s: float,
                       next_s: float | None, limit_s: int) -> None:
        """Still running at `elapsed_s`."""
        return None

    async def deliver(self, job: Job, outcome: Outcome) -> None:
        raise NotImplementedError

    async def succeeded(self, job: Job, outcome: Outcome) -> None:
        return None

    async def refund(self, job: Job, kind: str) -> int:
        raise NotImplementedError

    async def queued(self, job: Job, ahead: int) -> None:
        """Still waiting, with `ahead` jobs in front of it -- called when that changes."""
        return None

    async def card(self, job: Job) -> None:
        """Redraw the status while it runs; called every CARD_EVERY_S."""
        return None

    async def hold(self, job: Job) -> None:
        """A paid conversion has started: record it, so a hard kill cannot keep its credit."""
        return None

    async def release(self, job: Job) -> str:
        """It ended here: "held", "recovered" (already refunded elsewhere) or "absent"."""
        return "absent"

    async def failed(self, job: Job, kind: str, detail: str, refunded_to: int | None) -> None:
        return None

    async def alert(self, job: Job, kind: str, detail: str, elapsed: float,
                    refunded_to: int | None) -> None:
        return None

async def _safe(awaitable):
    try:
        return await awaitable
    except asyncio.CancelledError:
        raise
    except Exception:
        logger.exception("A conversion job hook failed")
        return None

def _unlink(path) -> None:
    try:
        Path(path).unlink(missing_ok=True)
    except OSError:
        logger.debug("Could not remove %s", path, exc_info=True)

async def _finish(job: Job, runner: Runner, kind: str | None, detail: str,
                  outcome: Outcome | None) -> None:
    try:
        elapsed = time.monotonic() - job.started_at if job.started_at else 0.0
        held = await _safe(runner.release(job)) if job.charged else None
        if kind is None:
            await _safe(runner.succeeded(job, outcome))
            return
        refunded_to = None
        if job.charged and kind in REFUNDED and held == "recovered":

            logger.warning("Conversion %s ended %s after it had been refunded as lost", job.id, kind)
            job.charged = False
        if job.charged and kind in REFUNDED:
            try:
                refunded_to = await runner.refund(job, kind)
                job.charged = False
            except Exception:
                logger.exception("Refunding conversion %s failed", job.id)
        await _safe(runner.failed(job, kind, detail, refunded_to))
        if kind in ALERTED:
            await _safe(runner.alert(job, kind, detail, elapsed, refunded_to))
    finally:
        _jobs.pop(job.id, None)
        job.state = "done"
        if outcome is not None and outcome.path:
            _unlink(outcome.path)
        if not job.keep_file:
            still_needed = active_paths()
            for path in job.paths:
                if str(path) not in still_needed:
                    _unlink(path)
        lock = _user_locks.get(job.user_id)
        if lock is not None and not lock.locked() and not getattr(lock, "_waiters", None):
            _user_locks.pop(job.user_id, None)

async def _watch_progress(job: Job, runner: Runner, limit: int) -> None:
    marks = [mark for mark in PROGRESS_MARKS_S if mark < limit]
    started = time.monotonic()
    for index, mark in enumerate(marks):
        await asyncio.sleep(max(0.0, mark - (time.monotonic() - started)))
        following = marks[index + 1] if index + 1 < len(marks) else None
        await _safe(runner.progress(job, index + 1, len(marks), mark, following, limit))

async def _watch_queue(job: Job, runner: Runner) -> None:
    shown = ahead_of(job)
    while True:
        await asyncio.sleep(CARD_EVERY_S)
        now = ahead_of(job)
        if now != shown:
            shown = now
            await _safe(runner.queued(job, now))

async def _watch_card(job: Job, runner: Runner) -> None:
    while True:
        await asyncio.sleep(CARD_EVERY_S)
        await _safe(runner.card(job))

async def _run(job: Job, runner: Runner) -> None:
    lock = _user_locks.setdefault(job.user_id, asyncio.Lock())
    kind: str | None = None
    detail = ""
    outcome: Outcome | None = None
    queue_watcher = asyncio.ensure_future(_watch_queue(job, runner))
    card_watcher = None
    try:
        async with lock:
            async with _global_slots():
                queue_watcher.cancel()
                if job.user_cancelled:
                    raise asyncio.CancelledError
                job.state = "running"
                job.started_at = time.monotonic()
                if await _safe(runner.refuse_start(job)):
                    kind, detail = PAUSED, "paused for an update before it could start"
                else:
                    balance = await runner.charge(job)
                    if balance is None:
                        kind = NO_CREDIT
                    else:
                        job.charged = job.price > 0
                        if job.charged:
                            await _safe(runner.hold(job))
                        job.extra.update(stage="converting", stage_started=time.monotonic(),
                                         progress=None, balance_after=balance)
                        await _safe(runner.started(job, balance))
                        limit = job_limit(job)
                        watcher = asyncio.ensure_future(_watch_progress(job, runner, limit))
                        card_watcher = asyncio.ensure_future(_watch_card(job, runner))
                        try:
                            outcome = await run_worker(job, runner.work_dir(),
                                                       f"{job.user_id}-{job.id}-out", limit)
                        finally:
                            watcher.cancel()
                        if not outcome.ok:
                            kind, detail = outcome.kind or CRASH, outcome.message
                        elif outcome.bytes > job.send_limit:
                            kind = TOO_LARGE
                            detail = f"{outcome.bytes / 1024 / 1024:.0f} MB"
                        else:
                            job.extra.update(stage="sending", stage_started=time.monotonic(), progress=None)
                            await _safe(runner.card(job))
                            try:
                                await runner.deliver(job, outcome)
                            except asyncio.CancelledError:
                                raise
                            except Exception as exc:
                                kind, detail = SEND_FAILED, f"{type(exc).__name__}: {exc}"
    except asyncio.CancelledError:
        queue_watcher.cancel()
        if card_watcher is not None:
            card_watcher.cancel()
        kind = USER_CANCELLED if job.user_cancelled else INTERRUPTED
        detail = ("stopped by the person who asked for it" if job.user_cancelled
                  else "the bot was stopped in the middle of it")
        await asyncio.shield(_finish(job, runner, kind, detail, outcome))
        if not job.user_cancelled:
            raise
        return
    except Exception as exc:
        logger.exception("Conversion job %s failed unexpectedly", job.id)
        kind, detail = CRASH, f"{type(exc).__name__}: {exc}"
    queue_watcher.cancel()
    if card_watcher is not None:
        card_watcher.cancel()
    job.extra["stage"] = "done"
    await _finish(job, runner, kind, detail, outcome)

def submit(application, job: Job, runner: Runner) -> Job:
    """Queue a job and start its task. Returns the job, whose `ahead_of` the caller can use to say where it stands."""
    _jobs[job.id] = job
    coroutine = _run(job, runner)
    if application is not None and hasattr(application, "create_task"):
        job.task = application.create_task(coroutine)
    else:
        job.task = asyncio.get_running_loop().create_task(coroutine)
    return job

def cancel(job_id: str, user_id: int | None = None) -> Job | None:
    """Stop a job at the request of the person it belongs to."""
    job = _jobs.get(job_id)
    if job is None or (user_id is not None and job.user_id != user_id):
        return None
    job.user_cancelled = True
    if job.task is not None and not job.task.done():
        job.task.cancel()
    return job

# ─── module: convert_bot.convert_runner ──────────────────────────────────────
"""The Telegram side of a conversion job."""
from __future__ import annotations

import asyncio
import logging
import os
import socket
import time
import uuid
from datetime import datetime, timezone
from pathlib import Path

from telegram import InlineKeyboardButton, InlineKeyboardMarkup

import big_files
import db
import family_link
import i18n
import jobs
import lifecycle
import live_message
import problems
from live_message import LiveMessage
from shared_features import emit_event, note_job, report_markup

logger = logging.getLogger(__name__)

BOT_LABEL = os.environ.get("FAMILY_LABEL") or "ConvertBot"

PROCESS_ID = f"{socket.gethostname()}:{os.getpid()}:{uuid.uuid4().hex[:6]}"

def _upper(ext: str) -> str:
    return (ext or "?").upper()

BAR_CELLS = 10

def progress_bar(fraction: float) -> str:
    """Ten cells and the percentage."""
    fraction = max(0.0, min(fraction, 0.99))
    filled = int(fraction * BAR_CELLS + 0.5)
    return "🟩" * filled + "⬜" * (BAR_CELLS - filled) + f" {int(fraction * 100)}%"

def clock(seconds: float) -> str:
    """0:07, 1:42, 12:05 -- a waiting time somebody can read at a glance."""
    seconds = max(0, int(seconds))
    return f"{seconds // 60}:{seconds % 60:02d}"

def time_left(done: float, total: float, elapsed: float) -> float | None:
    """Seconds still to go, from the rate so far -- or None until there is a rate worth believing: a guess from the first second is noise, and a number that swings wildly is worse than no number."""
    if not total or done <= 0 or elapsed < 3:
        return None
    fraction = done / total
    if fraction < 0.03 or fraction >= 1:
        return None
    return elapsed * (1 - fraction) / fraction

def status_card(lang: str, job: jobs.Job, now: float | None = None) -> str:
    now = time.monotonic() if now is None else now
    stage = job.extra.get("stage") or "converting"
    lines = [f"🔄 {_upper(job.src_ext)} → {_upper(job.target_ext)}",
             i18n.t(lang, "card_sending" if stage == "sending" else "card_converting")]
    progress = job.extra.get("progress")
    stage_started = job.extra.get("stage_started") or job.started_at or now
    timing = i18n.t(lang, "card_elapsed", elapsed=clock(now - (job.started_at or stage_started)))
    if progress and progress[1]:
        done, total = progress
        lines.append(progress_bar(done / total))
        left = time_left(done, total, now - stage_started)
        if left is not None:
            timing += " · " + i18n.t(lang, "card_left", left=clock(left))
    lines.append(timing)
    if job.price:
        lines.append(i18n.t(lang, "card_paid", price=job.price,
                            balance=job.extra.get("balance_after", "?")))
    else:
        lines.append(i18n.t(lang, "card_free"))
    return "\n".join(lines)

def download_card(lang: str, done: int, total: int, elapsed: float) -> str:
    """The same card for a file arriving, before there is a conversion."""
    lines = [i18n.t(lang, "card_downloading")]
    if total:
        lines.append(progress_bar(done / total))
        sizes = i18n.t(lang, "card_download_sizes", done=_megabytes(done), total=_megabytes(total))
        left = time_left(done, total, elapsed)
        if left is not None:
            sizes += " · " + i18n.t(lang, "card_left", left=clock(left))
        lines.append(sizes)
    return "\n".join(lines)

def _megabytes(size: int) -> str:
    return f"{size / 1024 / 1024:.1f} MB"

def result_filename(path) -> str:
    """What the converted file is called when it arrives."""
    stamp = datetime.now(timezone.utc).strftime("%Y-%m-%d_%H%M%S")
    return f"converted_{stamp}{Path(path).suffix}"

class ConvertRunner(jobs.Runner):
    def __init__(self, application, admin_ids, upload_dir, caption_for):
        self.application = application
        self.admin_ids = sorted(admin_ids or ())
        self.upload_dir = Path(upload_dir)
        self.caption_for = caption_for

    @property
    def bot(self):
        return self.application.bot

    def work_dir(self):
        self.upload_dir.mkdir(parents=True, exist_ok=True)
        return self.upload_dir

    def stop_keyboard(self, job: jobs.Job) -> InlineKeyboardMarkup:
        key = "job_stop_button" if job.state == "running" else "job_stop_button_queued"
        return InlineKeyboardMarkup([[InlineKeyboardButton(
            i18n.t(job.lang, key), callback_data=f"convstop:{job.id}")]])

    async def _status(self, job: jobs.Job, text: str, markup=None) -> None:
        """Rewrite the conversion's one status message."""
        live = job.extra.get("status_live")
        if live is None and job.extra.get("status_message") is not None:
            live = LiveMessage.adopt(job.extra["status_message"])
        try:
            if live is None:
                sent = await self.bot.send_message(chat_id=job.chat_id, text=text, reply_markup=markup)
                job.extra["status_message"] = sent
                job.extra["status_live"] = LiveMessage.adopt(sent)
                return
            before = live.message_id
            await live.set(self.bot, text, reply_markup=markup)
            job.extra["status_live"] = live
            if live.message_id != before:
                try:
                    await self.bot.delete_message(chat_id=live.chat_id, message_id=before)
                except Exception:
                    logger.debug("Could not remove the old status of conversion %s", job.id, exc_info=True)
        except Exception:
            logger.debug("Could not update the status of conversion %s", job.id, exc_info=True)

    async def refuse_start(self, job: jobs.Job) -> bool:
        return lifecycle.is_paused()

    async def charge(self, job: jobs.Job) -> int | None:
        if job.price <= 0:
            return await asyncio.to_thread(family_link.star_balance, job.user_id)
        return await asyncio.to_thread(
            family_link.spend_stars, job.user_id, job.price, "spend",
            f"{job.src_ext}->{job.target_ext} conversion, {job.size / 1024:.0f} KB")

    async def started(self, job: jobs.Job, balance: int) -> None:
        await self.card(job)

    async def queued(self, job: jobs.Job, ahead: int) -> None:
        if ahead:
            text = i18n.t(job.lang, "job_queued", src=_upper(job.src_ext), target=_upper(job.target_ext),
                          ahead=ahead, price=job.price)
        else:
            text = i18n.t(job.lang, "job_starting", src=_upper(job.src_ext), target=_upper(job.target_ext))
        await self._status(job, text, self.stop_keyboard(job))

    async def card(self, job: jobs.Job) -> None:
        """The live status, redrawn only when what it says has changed."""
        text = status_card(job.lang, job)
        if text == job.extra.get("card_text"):
            return
        job.extra["card_text"] = text
        stopping = job.extra.get("stage") != "sending"
        await self._status(job, text, self.stop_keyboard(job) if stopping else None)

    async def progress(self, job: jobs.Job, stage: int, stages: int, elapsed_s: float,
                       next_s, limit_s: int) -> None:
        """A new message, not an edit: an edit arrives silently, and the point is that somebody who has waited two minutes hears about it."""
        def minutes(seconds: float) -> int:
            return max(1, round(seconds / 60))

        if next_s is not None:
            text = i18n.t(job.lang, "job_progress", elapsed=minutes(elapsed_s),
                          next=minutes(next_s - elapsed_s), last=minutes(limit_s - elapsed_s))
        else:
            text = i18n.t(job.lang, "job_progress_last", elapsed=minutes(elapsed_s),
                          last=minutes(limit_s - elapsed_s))
        try:

            sent = await self.bot.send_message(chat_id=job.chat_id, text=text)
            live_message.bump(sent.chat_id, sent.message_id)
        except Exception:
            logger.debug("Could not send a progress update for conversion %s", job.id, exc_info=True)

    async def deliver(self, job: jobs.Job, outcome: jobs.Outcome) -> None:
        caption = self.caption_for(job.lang, outcome, job.src_ext, job.target_ext, job.price)
        if job.price:
            balance = await asyncio.to_thread(family_link.star_balance, job.user_id)
            caption += "\n\n" + i18n.t(job.lang, "balance_left_line", balance=balance)
        if outcome.bytes > jobs.SEND_LIMIT_BYTES and not jobs.LOCAL_BOT_API:

            def sending(done, total):
                job.extra["progress"] = (done, total)
            await big_files.send_document(job.chat_id, outcome.path,
                                          result_filename(outcome.path), caption, progress=sending)
            return

        with open(outcome.path, "rb") as handle:
            await self.bot.send_document(
                chat_id=job.chat_id, document=handle, filename=result_filename(outcome.path),
                caption=caption, read_timeout=300, write_timeout=300)

    async def succeeded(self, job: jobs.Job, outcome: jobs.Outcome) -> None:
        note_job(True)
        await self._status(job, i18n.t(job.lang, "job_done", src=_upper(job.src_ext),
                                       target=_upper(job.target_ext)))

    async def refund(self, job: jobs.Job, kind: str) -> int:
        return await asyncio.to_thread(
            family_link.move_stars, job.user_id, job.price, "refund",
            f"{job.src_ext}->{job.target_ext} conversion ended: {kind}")

    async def hold(self, job: jobs.Job) -> None:
        await asyncio.to_thread(db.hold_conversion, job.id, job.user_id, job.chat_id, job.price,
                                job.lang, f"{job.src_ext}->{job.target_ext}", PROCESS_ID)

    async def release(self, job: jobs.Job) -> str:
        return await asyncio.to_thread(db.release_conversion, job.id)

    async def keep_alive_and_recover(self) -> int:
        """Stamp this process's paid conversions, and give back the credit of any whose process was killed before it could."""
        await asyncio.to_thread(db.touch_conversions, PROCESS_ID)
        lost = await asyncio.to_thread(db.claim_orphaned_conversions, PROCESS_ID, jobs.HOLD_STALE_S)
        for row in lost:
            try:
                balance = await asyncio.to_thread(
                    family_link.move_stars, row["user_id"], row["price"], "refund",
                    f"{row['pair']} conversion ended: the bot was stopped before it finished")
            except Exception:
                logger.exception("Could not refund conversion %s, lost to a restart", row["job_id"])
                emit_event("error", "conversion",
                           f"⚠️ {BOT_LABEL}: a conversion lost to a restart could NOT be refunded "
                           f"({row['price']} ⚡, job {row['job_id']}). Refund it by hand.")
                continue
            lang = row["lang"] or "en"
            incident = problems.new_incident()
            text = (i18n.t(lang, "job_reason_interrupted") + "\n\n"
                    + i18n.t(lang, "job_refunded_line", price=row["price"], balance=balance)
                    + problems.code_line("CV-RESTARTED"))
            logger.info("Conversion %s was lost to a restart (CV-RESTARTED, incident %s); %s ⚡ refunded",
                        row["job_id"], incident, row["price"])
            try:
                await self.bot.send_message(chat_id=row["chat_id"], text=text,
                                            reply_markup=report_markup(lang, "CV-RESTARTED", incident))
            except Exception:
                logger.debug("Could not tell anyone about lost conversion %s", row["job_id"], exc_info=True)
            emit_event("warning", "conversion",
                       f"⚠️ {BOT_LABEL}: a conversion was lost to a restart and refunded "
                       f"({row['pair']}, {row['price']} ⚡ back). Incident {incident}.")
        return len(lost)

    async def failed(self, job: jobs.Job, kind: str, detail: str, refunded_to) -> None:
        lang = job.lang
        if kind == jobs.USER_CANCELLED:
            if job.started_at is None:
                text = i18n.t(lang, "job_removed_from_queue")
            elif job.price and job.charged:
                text = i18n.t(lang, "job_user_stopped", price=job.price)
            else:
                text = i18n.t(lang, "job_user_stopped_free")
            await self._status(job, text)
            return
        if kind == jobs.PAUSED:
            await self._status(job, i18n.t(lang, "update_soon_try_later_soon"))
            return
        if kind == jobs.NO_CREDIT:
            balance = await asyncio.to_thread(family_link.star_balance, job.user_id)
            await self._status(job, i18n.t(lang, "job_no_credit", price=job.price, balance=balance))
            return

        if kind in (jobs.TIMEOUT, jobs.CRASH, jobs.SEND_FAILED):
            note_job(False)
        limit = jobs.job_limit(job)
        reasons = {
            jobs.REFUSED: ("job_reason_refused", {"detail": detail}),
            jobs.TIMEOUT: ("job_reason_timeout", {"minutes": max(1, round(limit / 60))}),
            jobs.TOO_SLOW: ("job_reason_too_slow", {"minutes": max(2, round(job.extra.get("projected_s", 0) / 60)),
                                                    "limit": max(1, round(limit / 60))}),
            jobs.CRASH: ("job_reason_crash", {}),
            jobs.TOO_LARGE: ("job_reason_too_large", {"target": _upper(job.target_ext),
                                                      "mb": detail.split()[0] if detail else "?",
                                                      "limit": job.send_limit // (1024 * 1024)}),
            jobs.SEND_FAILED: ("job_reason_send_failed", {}),
            jobs.INTERRUPTED: ("job_reason_interrupted", {}),
        }
        key, fields = reasons.get(kind, ("job_reason_crash", {}))
        text = i18n.t(lang, key, **fields)
        if refunded_to is not None:
            text += "\n\n" + i18n.t(lang, "job_refunded_line", price=job.price, balance=refunded_to)

        code = problems.JOB_ENDINGS.get(kind, "CV-CRASH")
        incident = job.extra.setdefault("incident", problems.new_incident())
        logger.info("Conversion %s ended %s (%s, incident %s): %s", job.id, kind, code, incident, detail)
        await self._status(job, text + problems.code_line(code), report_markup(lang, code, incident))

    async def alert(self, job: jobs.Job, kind: str, detail: str, elapsed: float, refunded_to) -> None:
        """Straight to the owner, from this bot, in English."""
        limit = jobs.job_limit(job)

        credit = (f"{job.price} ⚡ back to the balance" if refunded_to is not None
                  else ("free conversion, nothing to refund" if not job.price else "not refunded"))
        text = (f"⚠️ {BOT_LABEL}: a conversion failed ({kind})\n"
                f"File: {job.size / 1024 / 1024:.1f} MB {_upper(job.src_ext)} → {_upper(job.target_ext)}\n"
                f"Ran: {elapsed:.0f}s of a {limit}s limit\n"
                f"Detail: {detail or '-'}\n"
                f"Code: {problems.JOB_ENDINGS.get(kind, 'CV-CRASH')} · incident {job.extra.get('incident', '-')}\n"
                f"Credit: {credit}")
        delivered = 0
        for admin_id in self.admin_ids:
            try:
                await self.bot.send_message(chat_id=admin_id, text=text)
                delivered += 1
            except Exception:
                logger.debug("Could not alert admin %s", admin_id, exc_info=True)
        emit_event("info" if delivered else "warning", "conversion", text)
